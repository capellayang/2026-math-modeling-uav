"""Run five identical Q3 searches and propagate fine-validated plans to strict Q4.

Only Q3AlgorithmConfig.seed differs between runs. Formal Q1–Q4 outputs are
read-only; all results are written to outputs/experiments/q3_five_seed/.
"""

from dataclasses import asdict, replace
from hashlib import sha256
import json
from pathlib import Path
import sys
import traceback

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from relief_uav.communication.coverage import RadioEnvironment
from relief_uav.data import load_scenario
from relief_uav.geo import build_segment_matrix
from relief_uav.geo.dem import DigitalElevationModel
from relief_uav.geo.segments import dem_source_path
from relief_uav.q3.model import Q3AlgorithmConfig
from relief_uav.q3.report import load_q3_solution
from relief_uav.q3.solver import solve_q3
from relief_uav.q4.model import Q4AlgorithmConfig
from relief_uav.q4.solver import solve_q4
from relief_uav.validation.q3 import validate_q3
from relief_uav.validation.q4 import validate_q4


SEEDS = (20260923, 20260924, 20260925, 20260926, 20260927)
CONFIG = Q3AlgorithmConfig()  # Branch defaults frozen; replace only seed below.
OUT = ROOT / "outputs" / "experiments" / "q3_five_seed"
BASELINE = ROOT / "outputs" / "q3" / "baseline_q2_v2_summary.json"
HEADERS = (
    "Seed", "Q3 status", "J1 weighted tardiness", "J2 joint makespan s",
    "J3 joint energy kWh", "NT transport sorties", "NR relay sorties",
    "strict components", "largest component service areas",
    "K2 legal partitions", "K3 legal partitions",
    "K2 Gap", "K2 Scale", "K2 CV",
    "K3 Gap", "K3 Scale", "K3 CV",
    "0.25s validator", "OUTAGE s", "Q4 validator",
    "Q3 5-axis Pareto", "Q3 near-Pareto", "final choice",
    "Q3 search seconds", "ALNS iterations completed", "note",
)


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2,
                               allow_nan=False), encoding="utf-8")


def _source_fingerprint() -> str:
    digest = sha256()
    for path in sorted((ROOT / "src" / "relief_uav" / "q3").glob("*.py")):
        digest.update(path.name.encode())
        digest.update(path.read_bytes())
    return digest.hexdigest()


def _strict_metrics(q4, k: int) -> dict:
    result = next(p for p in q4.policy_results
                  if p.relay_policy == "strict_no_duplication" and p.k == k)
    selected = next((c for c in result.candidates if c.selected), None)
    return {
        "legal_partitions": len(result.candidates),
        "gap": selected.stock_gap_total if selected else None,
        "scale": selected.resource_scale if selected else None,
        "cv": selected.workload_cv if selected else None,
    }


def _rank(rows: list[dict]) -> None:
    for row in rows:
        row["pareto"] = False
        row["near_pareto"] = False
        row["final_choice"] = False
    valid = [r for r in rows if r["status"] == "PASS"]
    for row in valid:
        metrics = (row["j1"], row["j2_s"], row["j3_kwh"], row["nt"], row["nr"])
        row["pareto"] = not any(
            other is not row and
            all(a <= b + 1e-8 for a, b in zip(
                (other["j1"], other["j2_s"], other["j3_kwh"], other["nt"], other["nr"]),
                metrics)) and
            any(a < b - 1e-8 for a, b in zip(
                (other["j1"], other["j2_s"], other["j3_kwh"], other["nt"], other["nr"]),
                metrics))
            for other in valid)
    if not valid:
        return
    best_j2 = min(r["j2_s"] for r in valid)
    best_j3 = min(r["j3_kwh"] for r in valid)
    for row in valid:
        row["near_pareto"] = (abs(row["j1"]) <= 1e-8 and
                               row["j2_s"] <= 1.02 * best_j2 + 1e-8 and
                               row["j3_kwh"] <= 1.01 * best_j3 + 1e-8)
    eligible = [r for r in valid if r["near_pareto"]]
    if eligible:
        # First preserve the Q3 quality gates above, then prefer a partitionable
        # communication structure. Missing K3 is worse than any feasible K3.
        chosen = min(eligible, key=lambda r: (
            -r["components"], r["largest_component"],
            -r["k3"]["legal_partitions"], -r["k2"]["legal_partitions"],
            r["k3"]["gap"] if r["k3"]["gap"] is not None else float("inf"),
            r["k3"]["cv"] if r["k3"]["cv"] is not None else float("inf"),
            r["k2"]["gap"] if r["k2"]["gap"] is not None else float("inf"),
            r["k2"]["cv"] if r["k2"]["cv"] is not None else float("inf"),
            r["j2_s"], r["j3_kwh"], r["seed"]))
        chosen["final_choice"] = True


def _save_workbook(rows: list[dict]) -> None:
    book = Workbook()
    ws = book.active
    ws.title = "Five seeds"
    ws.append(HEADERS)
    for r in rows:
        k2, k3 = r.get("k2", {}), r.get("k3", {})
        ws.append((r["seed"], r["status"], r.get("j1"), r.get("j2_s"),
                   r.get("j3_kwh"), r.get("nt"), r.get("nr"),
                   r.get("components"), r.get("largest_component"),
                   k2.get("legal_partitions"), k3.get("legal_partitions"),
                   k2.get("gap"), k2.get("scale"), k2.get("cv"),
                   k3.get("gap"), k3.get("scale"), k3.get("cv"),
                   r.get("fine_status"), r.get("outage_s"), r.get("q4_status"),
                   r.get("pareto"), r.get("near_pareto"),
                   bool(r.get("final_choice")), r.get("search_seconds"),
                   r.get("alns_iterations"), r.get("note")))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24445C")
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(66, max(15,
            max(len(str(c.value or "")) for c in column) + 2))
    config_ws = book.create_sheet("Fixed configuration")
    config_ws.append(("parameter", "value"))
    for key, value in asdict(CONFIG).items():
        config_ws.append((key, json.dumps(value, ensure_ascii=False)))
    config_ws.append(("seeds", json.dumps(SEEDS)))
    config_ws.append(("q3_source_sha256", _source_fingerprint()))
    config_ws.append(("q2_baseline_sha256", sha256(BASELINE.read_bytes()).hexdigest()))
    config_ws.append(("fine_validation_step_s", 0.25))
    config_ws.append(("q4_policy", "strict"))
    config_ws.append(("near_pareto", "J1=0; J2<=1.02*best_valid_J2; J3<=1.01*best_valid_J3"))
    config_ws.column_dimensions["A"].width = 35
    config_ws.column_dimensions["B"].width = 90
    book.save(OUT / "q3_five_seed_comparison.xlsx")


def _checkpoint(rows: list[dict]) -> None:
    _rank(rows)
    _write_json(OUT / "comparison.json", {"fixed_config": asdict(CONFIG),
        "seeds": SEEDS, "q3_source_sha256": _source_fingerprint(), "rows": rows})
    _save_workbook(rows)


def _run_seed(seed: int, scenario, segments, env) -> dict:
    row = {"seed": seed, "status": "FAIL", "fine_status": "NOT_RUN",
           "q4_status": "NOT_RUN", "pareto": False, "near_pareto": False,
           "final_choice": False}
    seed_dir = OUT / str(seed)
    seed_dir.mkdir(parents=True, exist_ok=True)
    config = replace(CONFIG, seed=seed)
    run = solve_q3(scenario, segments, env, config, baseline_path=BASELINE,
                   experiment_cache_dir=seed_dir / "cache")
    source_path = seed_dir / "q3.json"
    _write_json(source_path, asdict(run.selected))
    # Re-read the saved artifact and reconstruct radio/DEM independently from
    # the optimizer, even though solve_q3 itself also performed a 0.25 s check.
    saved = load_q3_solution(source_path)
    independent = RadioEnvironment(scenario, DigitalElevationModel(dem_source_path(ROOT)))
    fine = validate_q3(independent, segments, saved, max_step_s=0.25,
                       transition_tolerance_s=config.transition_tolerance_s)
    row.update({"fine_status": "PASS" if fine.passed else "FAIL",
                "outage_s": fine.outage_s, "search_seconds": run.elapsed_s,
                "alns_iterations": saved.alns_iterations,
                "j1": saved.objective.weighted_tardiness,
                "j2_s": saved.objective.joint_makespan_s,
                "j3_kwh": saved.objective.joint_energy_kwh,
                "nt": saved.objective.transport_sortie_count,
                "nr": saved.objective.relay_sortie_count})
    _write_json(seed_dir / "q3_fine_validation.json", {
        "passed": fine.passed, "outage_s": fine.outage_s,
        "max_step_s": fine.verification_max_step_s,
        "issues": fine.issues})
    if not fine.passed or abs(fine.outage_s) > 1e-9:
        row["note"] = "; ".join(str(x) for x in fine.issues[:5]) or "OUTAGE > 0"
        return row
    q4 = solve_q4(scenario, saved, source_path,
                  {"passed": True, "outage_s": fine.outage_s},
                  Q4AlgorithmConfig(relay_policy="strict"))
    q4_payload = json.loads(json.dumps(asdict(q4), ensure_ascii=False))
    q4_path = seed_dir / "q4_strict.json"
    _write_json(q4_path, q4_payload)
    q4_check = validate_q4(scenario, saved, source_path, q4_payload, q3_passed=True)
    row["q4_status"] = "PASS" if q4_check.passed else "FAIL"
    _write_json(seed_dir / "q4_validation.json", asdict(q4_check))
    if not q4_check.passed:
        row["note"] = "; ".join(q4_check.issues[:5])
        return row
    row.update({"status": "PASS", "components": len(q4.strict_components),
                "largest_component": max(len(c.service_areas) for c in q4.strict_components),
                "k2": _strict_metrics(q4, 2), "k3": _strict_metrics(q4, 3),
                "note": "0.25s Q3 and independent exact Q4 PASS"})
    return row


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    scenario = load_scenario(ROOT)
    segments = build_segment_matrix(scenario)
    env = RadioEnvironment(scenario, DigitalElevationModel(dem_source_path(ROOT)))
    rows = []
    for seed in SEEDS:
        print(f"START seed={seed}", flush=True)
        try:
            row = _run_seed(seed, scenario, segments, env)
        except Exception as exc:
            row = {"seed": seed, "status": "FAIL", "fine_status": "NOT_RUN",
                   "q4_status": "NOT_RUN", "pareto": False,
                   "near_pareto": False, "final_choice": False,
                   "note": f"{type(exc).__name__}: {exc}"}
            (OUT / str(seed)).mkdir(parents=True, exist_ok=True)
            (OUT / str(seed) / "error.txt").write_text(
                traceback.format_exc(), encoding="utf-8")
        rows.append(row)
        _checkpoint(rows)
        print(f"DONE seed={seed} status={row['status']} "
              f"fine={row.get('fine_status')} Q4={row.get('q4_status')} "
              f"components={row.get('components')}", flush=True)
    print(f"Comparison: {OUT / 'q3_five_seed_comparison.xlsx'}", flush=True)
    if not all(r["status"] == "PASS" for r in rows):
        raise SystemExit("One or more seeds failed; inspect comparison.json and error.txt")


if __name__ == "__main__":
    main()
