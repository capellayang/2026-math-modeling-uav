"""Reproducible, output-isolated Q3/Q4 sensitivity experiments.

Run with the project's Miniconda dl interpreter from the repository root.
All snapshots, tables and figures stay under outputs/experiments/q3_q4_improvement.
"""

from dataclasses import asdict, replace
import json
from pathlib import Path
import sys
from time import monotonic

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from relief_uav.communication.coverage import RadioEnvironment, coverage_intervals
from relief_uav.data import load_scenario
from relief_uav.geo import build_segment_matrix
from relief_uav.geo.dem import DigitalElevationModel
from relief_uav.geo.segments import dem_source_path
from relief_uav.q2.report import load_q2_solution
from relief_uav.q3.candidates import generate_candidates
from relief_uav.q3.joint_cp import schedule_joint
from relief_uav.q3.model import Q3AlgorithmConfig
from relief_uav.q3.report import load_q3_solution
from relief_uav.q3.relay_physics import construct_relay_sortie
from relief_uav.q3.timeline import build_q3_solution
from relief_uav.q4.inheritance import atomic_components
from relief_uav.q4.model import Q4AlgorithmConfig
from relief_uav.q4.partition import enumerate_partitions, canonical_signature
from relief_uav.q4.resources import global_minimum, group_metrics, stock_vector
from relief_uav.q4.solver import build_task_groups, solve_q4
from relief_uav.validation.q2 import validate_q2
from relief_uav.validation.q3 import validate_q3
from relief_uav.validation.q4 import validate_q4


OUT = ROOT / "outputs/experiments/q3_q4_improvement"
CACHE = OUT / "cache"
FIG = OUT / "figures"


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2,
                               allow_nan=False), encoding="utf-8")


def sheet(name, headers, rows):
    OUT.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    ws = book.active
    ws.title = "Results"
    ws.append(headers)
    for row in rows:
        ws.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (tuple, list, dict))
                   else v for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24445C")
    for index, col in enumerate(ws.columns, 1):
        width = min(70, max(16, max(len(str(c.value or "")) for c in col[:100])+2))
        ws.column_dimensions[get_column_letter(index)].width = width
    book.save(OUT/name)


def q4_rows(scenario, q3, source, validation):
    result = solve_q4(scenario, q3, source,
                      {"passed": validation.passed, "outage_s": validation.outage_s},
                      Q4AlgorithmConfig(relay_policy="both"))
    # The saved Q4 schema uses JSON arrays, whereas dataclasses contain tuples.
    payload = json.loads(json.dumps(asdict(result), ensure_ascii=False))
    check = validate_q4(scenario, q3, source, payload, q3_passed=validation.passed)
    if not check.passed:
        raise RuntimeError(f"Q4 validator FAIL: {check.issues[:5]}")
    return result


def q4_key(result, policy, k):
    p = next(x for x in result.policy_results if x.relay_policy == policy and x.k == k)
    if not p.feasible:
        return {"components": p.atomic_component_count, "gap": None,
                "scale": None, "cv": None, "candidate": None}
    chosen = next(c for c in p.candidates if c.selected)
    return {"components": p.atomic_component_count, "gap": chosen.stock_gap_total,
            "scale": chosen.resource_scale, "cv": chosen.workload_cv,
            "candidate": chosen.candidate_id}


def selection_rows(name, q4):
    rows = []
    for policy in ("strict_no_duplication", "replicate_relay"):
      for k in (2, 3):
        p = next(x for x in q4.policy_results if x.relay_policy == policy and x.k == k)
        if not p.feasible:
            continue
        candidates = p.candidates
        min_gap = min(c.stock_gap_total for c in candidates)
        for dg in range(0, min(8, max(c.stock_gap_total for c in candidates)-min_gap)+1):
            gap_ok = [c for c in candidates if c.stock_gap_total <= min_gap+dg]
            min_scale = min(c.resource_scale for c in gap_ok)
            for ds in range(0, min(8, max(c.resource_scale for c in gap_ok)-min_scale)+1):
                eligible = [c for c in gap_ok if c.resource_scale <= min_scale+ds]
                chosen = min(eligible, key=lambda c: (c.workload_cv,
                                c.stock_gap_total, c.resource_scale,
                                c.canonical_group_signature))
                rows.append((name, policy, k, dg, ds, chosen.candidate_id,
                             chosen.stock_gap_total, chosen.resource_scale,
                             chosen.workload_cv, chosen.canonical_group_signature))
        # Normalized distance from the three-dimensional ideal on the exact
        # nondominated set is an explicitly labeled experimental knee rule.
        front = [c for c in candidates if c.pareto]
        lo = [min(getattr(c, key) for c in candidates) for key in
              ("stock_gap_total", "resource_scale", "workload_cv")]
        hi = [max(getattr(c, key) for c in candidates) for key in
              ("stock_gap_total", "resource_scale", "workload_cv")]
        def ideal_distance(c):
            vals = (c.stock_gap_total, c.resource_scale, c.workload_cv)
            return sum(((v-a)/(b-a) if b>a else 0)**2
                       for v, a, b in zip(vals, lo, hi))**.5
        knee = min(front, key=lambda c: (ideal_distance(c), c.canonical_group_signature))
        rows.append((name, policy, k, "KNEE", "KNEE", knee.candidate_id,
                     knee.stock_gap_total, knee.resource_scale,
                     knee.workload_cv, knee.canonical_group_signature))
    return rows


def workload_rows(name, q3, q4):
    rows = []
    for p in q4.policy_results:
        if not p.feasible:
            continue
        c = next(x for x in p.candidates if x.selected)
        for g in c.groups:
            relay = {r.sortie_id: r for r in q3.relays}
            service = sum(relay[t["source_sortie_id"]].service_end_s-
                          relay[t["source_sortie_id"]].service_start_s
                          for t in g.group.relay_tasks)
            resource_hours = {key: sum(a.occupied_end_s-a.occupied_start_s
                                      for a in g.assignments if a.resource_type == key)/3600
                              for key in ("A_UAV", "B_UAV", "C_UAV", "A_battery",
                                          "B_battery", "C_battery", "relay_UAV",
                                          "relay_component")}
            rows.append((name, p.relay_policy, p.k, g.group.group_id,
                         g.service_area_count, g.box_count, g.cargo_mass_kg,
                         g.transport_sortie_count, g.transport_workload_s,
                         service, g.relay_workload_s, g.total_workload_s,
                         sum(resource_hours[k] for k in ("A_UAV", "B_UAV", "C_UAV")),
                         sum(resource_hours[k] for k in ("A_battery", "B_battery", "C_battery")),
                         resource_hours["relay_UAV"], resource_hours["relay_component"]))
    return rows


def group_local_policy(scenario, env, q3, k, *, validation_step_s=4.0,
                       delta_gap=0, delta_scale=0):
    """Experimental interpretation: trim cross-group relay service windows.

    Each partition receives its own relay resources. Q3 transport route/box/start
    and original hover/AGL stay frozen. This is intentionally separate from the
    official Q4 policy and has its own communication/resource validation.
    """
    atoms = atomic_components(scenario, q3, strict=False)
    original_relays = {r.sortie_id: r for r in q3.relays}
    by_transport = {s.spec.sortie_id: s for s in q3.transport.sorties}
    stock = stock_vector(scenario)
    minimum = global_minimum(scenario, q3)
    rows = []
    for idx, partition in enumerate(enumerate_partitions(atoms, k), 1):
        service_group = {sid: g for g, sites in enumerate(partition) for sid in sites}
        transport_group = {}
        for tid, sortie in by_transport.items():
            labels = {service_group[s] for s in sortie.spec.route if s != "O01"}
            if len(labels) != 1:
                raise AssertionError("Transport atom split")
            transport_group[tid] = labels.pop()
        required = {}
        for interval in q3.communication:
            if interval.mode == "RELAY":
                key = (interval.relay_sortie_id,
                       transport_group[interval.transport_sortie_id])
                required.setdefault(key, []).append(interval)
        copies, remap = [], {}
        for (rid, group), intervals in sorted(required.items()):
            original = original_relays[rid]
            lo = max(original.service_start_s, min(r.start_s for r in intervals)-.05)
            hi = min(original.service_end_s, max(r.end_s for r in intervals)+.05)
            new_id = f"GL-K{k}-P{idx:03d}-G{group+1}-{rid}"
            try:
                copy = construct_relay_sortie(scenario, env.dem, new_id,
                        original.drone_id, original.component_id, original.hover,
                        lo, hi, setup_energy_mode=q3.setup_energy_mode)
            except ValueError:
                copies = []
                break
            copies.append(copy)
            remap[(rid, group)] = new_id
        if not copies:
            continue
        communication = tuple(replace(c, relay_sortie_id=remap[
            (c.relay_sortie_id, transport_group[c.transport_sortie_id])])
            if c.mode == "RELAY" else c for c in q3.communication)
        local = replace(q3, relays=tuple(copies), communication=communication)
        groups = build_task_groups(local, partition, k, "strict_no_duplication")
        metrics = tuple(group_metrics(scenario, local, g) for g in groups)
        vector = {key: sum(g.minimum_recolored_requirement[key] for g in metrics)
                  for key in stock}
        gap = sum(max(0, vector[key]-stock[key]) for key in stock)
        scale = sum(vector.values())
        work = [g.total_workload_s for g in metrics]
        mean = sum(work)/k
        cv = (sum((w-mean)**2 for w in work)/k)**.5/mean
        rows.append((gap, scale, cv, canonical_signature(partition),
                     local, metrics, vector))
    if not rows:
        return None
    min_gap = min(r[0] for r in rows)
    gap_eligible = [r for r in rows if r[0] <= min_gap+delta_gap]
    min_scale = min(r[1] for r in gap_eligible)
    eligible = [r for r in gap_eligible if r[1] <= min_scale+delta_scale]
    chosen = min(eligible, key=lambda r: (r[2], r[0], r[1], r[3]))
    gap, scale, cv, signature, local, metrics, vector = chosen
    # Verify every frozen transport trajectory with only its local group's
    # freshly calculated relay missions, at the same fine step as Q3.
    issues = []
    for g in metrics:
        relays = tuple(r for r in local.relays
                       if any(t["source_sortie_id"] in r.sortie_id
                              for t in g.group.relay_tasks))
        for tid in g.group.transport_sorties:
            intervals, _ = coverage_intervals(env, by_transport[tid], relays,
                                              max_step_s=validation_step_s)
            if any(x.mode == "OUTAGE" for x in intervals):
                issues.append(f"{g.group.group_id}/{tid} OUTAGE")
    return {"k": k, "delta_gap": delta_gap, "delta_scale": delta_scale,
            "gap": gap, "scale": scale, "cv": cv,
            "partition": signature, "relay_sorties": len(local.relays),
            "relay_energy_kwh": sum(r.total_energy_kwh for r in local.relays),
            "relay_components": vector["relay_component"],
            "global_minimum": minimum,
            "validator": "PASS" if not issues else "FAIL",
            "validation_max_step_s": validation_step_s,
            "issues": issues, "enumerated": len(rows),
            "candidates": [{"gap": r[0], "scale": r[1], "cv": r[2],
                "partition": r[3], "relay_sorties": len(r[4].relays),
                "relay_energy_kwh": sum(x.total_energy_kwh for x in r[4].relays),
                "relay_components": r[6]["relay_component"],
                "selected": r is chosen} for r in rows]}


def profile(scenario, segments, env, transport, audit, name, config,
            *, budget_s=30):
    stamp = monotonic()
    cache_path = CACHE / (f"hover_{config.hover_xy_mode}_{config.hover_altitude_mode}_"
                          f"{int(config.hover_grid_m)}_{config.hover_top_k}.json")
    hover, representatives, hit = generate_candidates(
        env, transport, audit, config, cache_path)
    print(name, "hover", len(hover), "representatives", len(representatives),
          "cache_hit", hit, "seconds", round(monotonic()-stamp, 1), flush=True)
    if not hover:
        return None, {"name": name, "status": "NO_HOVER", "hover_count": 0}
    try:
        cp = schedule_joint(env, segments, transport, audit, hover, config,
                            time_limit_s=budget_s)
    except ValueError as exc:
        return None, {"name": name, "status": f"MODEL_ERROR: {exc}",
                      "hover_count": len(hover)}
    if cp is None:
        return None, {"name": name, "status": "NO_CP_INCUMBENT",
                      "hover_count": len(hover)}
    q3 = build_q3_solution(env, cp.transport, cp.relays, seed=config.seed,
                           search_seconds=monotonic()-stamp,
                           cp_sat_status=cp.cp_sat_status, coverage_step_s=None)
    check = validate_q3(env, segments, q3, max_step_s=4.0)
    meta = {"name": name, "status": "PASS" if check.passed else "FAIL",
            "issues": check.issues[:5], "hover_count": len(hover),
            "hover_heights": sorted({p.hover_agl_m for p in hover}),
            "cp_status": cp.cp_sat_status, "elapsed_s": monotonic()-stamp,
            "max_relay_sorties": config.max_relay_sorties,
            "xy_mode": config.hover_xy_mode,
            "altitude_mode": config.hover_altitude_mode,
            "tardiness_mode": config.tardiness_mode,
            "epsilon": config.absolute_epsilon,
            "joint_objective": config.joint_objective}
    meta["validation_max_step_s"] = 4.0
    if not check.passed:
        return None, meta
    q3 = replace(q3, communication=check.intervals)
    return (q3, check), meta


def plot(rows, q4_comparisons):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    FIG.mkdir(parents=True, exist_ok=True)
    passed = [r for r in rows if r["status"] == "PASS"]
    for filename, xkey, ykey in (
        ("q3_pareto_extended.png", "j1", "j2"),
        ("q3_vs_q4_tradeoff.png", "j2", "k2_strict_cv"),
        ("q4_k2_tradeoff.png", "k2_strict_gap", "k2_strict_cv"),
        ("q4_k3_tradeoff.png", "k3_strict_gap", "k3_strict_cv")):
        fig, ax = plt.subplots(figsize=(8, 5))
        for row in passed:
            x, y = row.get(xkey), row.get(ykey)
            if x is None or y is None:
                continue
            ax.scatter(x, y)
            ax.annotate(row["name"], (x, y), fontsize=8)
        ax.set_xlabel(xkey)
        ax.set_ylabel(ykey)
        ax.grid(alpha=.25)
        fig.tight_layout()
        fig.savefig(FIG/filename, dpi=160)
        plt.close(fig)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    scenario = load_scenario(ROOT)
    segments = build_segment_matrix(scenario)
    env = RadioEnvironment(scenario, DigitalElevationModel(dem_source_path(ROOT)))
    transport = load_q2_solution(ROOT/"outputs/q3/baseline_q2_v2_summary.json")
    if not validate_q2(scenario, segments, transport).passed:
        raise RuntimeError("Baseline Q2 validator FAIL")
    audit = json.loads((ROOT/"outputs/q3/baseline_q2_direct_audit.json").read_text(encoding="utf-8"))
    formal = ROOT/"outputs/q3/q3_summary.json"
    current = load_q3_solution(formal)
    current_check = validate_q3(env, segments, current, max_step_s=4.0)
    if not current_check.passed:
        raise RuntimeError(f"Formal Q3 validator FAIL: {current_check.issues[:5]}")
    results = [("CURRENT", current, current_check, formal)]
    metadata = [{"name": "CURRENT", "status": "PASS", "hover_count": 20,
                 "altitude_mode": "legacy", "xy_mode": "local", "max_relay_sorties": 6}]
    base = Q3AlgorithmConfig()
    trials = [
        ("COMPONENT_REUSE_6", replace(base, max_relay_sorties=6)),
        ("COMPONENT_REUSE_8", replace(base, max_relay_sorties=8)),
        ("COMPONENT_REUSE_10", replace(base, max_relay_sorties=10)),
        ("COMPONENT_REUSE_12", replace(base, max_relay_sorties=12)),
        ("ALTITUDE_SEARCH", replace(base, hover_altitude_mode="full")),
        ("XY_BUFFERED", replace(base, hover_xy_mode="buffered", hover_grid_m=800)),
        ("XY_EXPANDED", replace(base, hover_xy_mode="full_dem", hover_grid_m=800)),
        ("MULTIOBJECTIVE_0", replace(base, tardiness_mode="pareto",
                                     route_archive_mode="multiobjective",
                                     max_relay_sorties=8, absolute_epsilon=0)),
        ("MULTIOBJECTIVE_1000", replace(base, tardiness_mode="pareto",
                                        route_archive_mode="multiobjective",
                                        max_relay_sorties=8, absolute_epsilon=1000)),
        ("MULTIOBJECTIVE_5000", replace(base, tardiness_mode="pareto",
                                        route_archive_mode="multiobjective",
                                        max_relay_sorties=8, absolute_epsilon=5000)),
        ("MULTIOBJECTIVE_ENERGY", replace(base, tardiness_mode="pareto",
                                        route_archive_mode="multiobjective",
                                        max_relay_sorties=8, absolute_epsilon=5000,
                                        joint_objective="energy")),
        ("MULTIOBJECTIVE_SORTIES", replace(base, tardiness_mode="pareto",
                                        route_archive_mode="multiobjective",
                                        max_relay_sorties=8, absolute_epsilon=5000,
                                        joint_objective="relay_sorties")),
    ]
    for name, config in trials:
        previous_path = OUT/f"{name.lower()}_q3.json"
        if previous_path.is_file():
            previous = load_q3_solution(previous_path)
            previous_check = validate_q3(env, segments, previous, max_step_s=4)
            if previous_check.passed:
                previous_meta = {"name": name, "status": "PASS",
                    "hover_count": None, "max_relay_sorties": config.max_relay_sorties,
                    "xy_mode": config.hover_xy_mode,
                    "altitude_mode": config.hover_altitude_mode,
                    "tardiness_mode": config.tardiness_mode,
                    "epsilon": config.absolute_epsilon,
                    "joint_objective": config.joint_objective,
                    "validation_max_step_s": 4.0, "resumed": True}
                metadata.append(previous_meta)
                results.append((name, previous, previous_check, previous_path))
                print(name, "resumed PASS", flush=True)
                continue
        found, meta = profile(scenario, segments, env, transport, audit,
                              name, config, budget_s=30)
        metadata.append(meta)
        print(name, meta["status"], round(meta.get("elapsed_s", 0), 1), flush=True)
        if found:
            q3, check = found
            path = OUT/f"{name.lower()}_q3.json"
            save_json(path, asdict(q3))
            results.append((name, q3, check, path))
        save_json(OUT/"experiment_progress.json", metadata)
    q3_rows, q4_rows_out, selection, workload = [], [], [], []
    local_rows = []
    for name, q3, check, source in results:
        q4 = q4_rows(scenario, q3, source, check)
        save_json(OUT/f"{name.lower()}_q4.json", asdict(q4))
        selection.extend(selection_rows(name, q4))
        workload.extend(workload_rows(name, q3, q4))
        o = q3.objective
        row = {"name": name, "status": "PASS", "j1": o.weighted_tardiness,
               "j2": o.joint_makespan_s, "j3": o.joint_energy_kwh,
               "transport_sorties": o.transport_sortie_count,
               "relay_sorties": o.relay_sortie_count,
               "min_access_db": check.minimum_relay_access_margin_db,
               "min_backhaul_db": check.minimum_relay_backhaul_margin_db,
               "outage_s": check.outage_s}
        for k in (2, 3):
            for policy, label in (("strict_no_duplication", "strict"),
                                  ("replicate_relay", "replicate")):
                metrics = q4_key(q4, policy, k)
                for key, value in metrics.items():
                    row[f"k{k}_{label}_{key}"] = value
                q4_rows_out.append((name, policy, k, metrics["components"],
                                    metrics["gap"], metrics["scale"], metrics["cv"],
                                    metrics["candidate"], "PASS"))
        q3_rows.append(row)
        save_json(OUT/"comparison_progress.json", q3_rows)
    combined = [{**m, **next((r for r in q3_rows if r["name"] == m["name"]), {})}
                for m in metadata]
    q3_headers = ["name", "status", "j1", "j2", "j3", "transport_sorties",
                  "relay_sorties", "min_access_db", "min_backhaul_db", "outage_s",
                  "k2_strict_components", "k2_strict_gap", "k2_strict_scale",
                  "k2_strict_cv", "k3_strict_components", "k3_strict_gap",
                  "k3_strict_scale", "k3_strict_cv", "issues"]
    sheet("q3_candidate_comparison.xlsx", q3_headers,
          [[r.get(k) for k in q3_headers] for r in combined])
    for filename, selected in (
        ("q3_component_reuse_comparison.xlsx", [r for r in combined if r["name"] == "CURRENT" or r["name"].startswith("COMPONENT_REUSE")]),
        ("q3_hover_search_comparison.xlsx", [r for r in combined if r["name"] in
            ("CURRENT", "ALTITUDE_SEARCH", "XY_BUFFERED", "XY_EXPANDED")]),
        ("q3_multiobjective_pareto.xlsx", [r for r in combined if r["name"] == "CURRENT" or r["name"].startswith("MULTIOBJECTIVE")])):
        sheet(filename, q3_headers+["hover_count", "hover_heights", "elapsed_s",
                                  "max_relay_sorties", "altitude_mode", "xy_mode",
                                  "tardiness_mode", "epsilon", "joint_objective"],
              [[r.get(k) for k in q3_headers+["hover_count", "hover_heights", "elapsed_s",
                                          "max_relay_sorties", "altitude_mode", "xy_mode",
                                          "tardiness_mode", "epsilon", "joint_objective"]]
               for r in selected])
    sheet("q4_cross_q3_comparison.xlsx",
          ["q3", "policy", "K", "components", "gap", "scale", "CV",
           "selected_candidate", "validator"], q4_rows_out)
    sheet("q4_selection_sensitivity.xlsx",
          ["q3", "policy", "K", "delta_gap", "delta_scale", "candidate",
           "gap", "scale", "CV", "partition"], selection)
    sheet("q4_workload_sensitivity.xlsx",
          ["q3", "policy", "K", "group", "service_areas", "boxes", "cargo_kg",
           "transport_sorties", "transport_operation_s", "relay_service_s",
           "relay_mission_s", "project_workload_s", "transport_UAV_hours",
           "transport_battery_hours", "relay_UAV_hours", "relay_component_hours"], workload)
    sheet("q4_relay_policy_comparison.xlsx",
          ["q3", "policy", "K", "components", "gap", "scale", "CV",
           "selected_candidate", "validator", "relay_sorties", "relay_energy_kwh",
           "relay_components", "enumerated"],
          [row+(None, None, None, None) for row in q4_rows_out]+local_rows)
    plot(q3_rows, q4_rows_out)
    print("complete", OUT, flush=True)


if __name__ == "__main__":
    main()
