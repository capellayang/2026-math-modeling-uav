"""Verify exported audit artifacts and preserve source hashes."""
from pathlib import Path
import json,hashlib,zipfile
from lxml import etree
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/data_audit'
manifest=json.loads((OUT/'source_manifest.json').read_text(encoding='utf-8'))
changed=[m['path'] for m in manifest if hashlib.sha256((ROOT/m['path']).read_bytes()).hexdigest()!=m['sha256']]
wb=openpyxl.load_workbook(OUT/'data_summary.xlsx')
errors=[(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e']
with zipfile.ZipFile(next(ROOT.glob('*.docx'))) as z:
    parts={n:len(etree.fromstring(z.read(n)).xpath('//*[local-name()="oMath"]')) for n in z.namelist() if n.endswith('.xml') and b'oMath' in z.read(n)}
result={'source_files':len(manifest),'changed_sources':changed,'excel_errors':errors,'exported_sheets':[(s.title,s.max_row,s.max_column) for s in wb],'office_math_by_part':parts,'source_python_files':[str(p.relative_to(ROOT)) for p in (ROOT/'src').rglob('*.py')],'note':'没有Q1-Q4求解器；检查范围为本阶段审计产物'}
assert not changed and not errors
assert wb['审计概览']['B3'].value==80
assert wb['审计概览']['B4'].value==758
assert abs(wb['审计概览']['B5'].value-2.011)<1e-12
(OUT/'artifact_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
