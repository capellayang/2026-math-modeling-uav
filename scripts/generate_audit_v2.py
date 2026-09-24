"""Source-to-code audit tables; rules are labeled by provenance and risk."""

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/"scripts"))
from q3_q4_improvement import sheet

OUT = ROOT/"outputs/audit_v2"
OUT.mkdir(parents=True, exist_ok=True)


def workbook(name, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    w = Workbook(); s = w.active; s.title = "审计"
    s.append(headers)
    for row in rows:
        s.append(row)
    s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="24445C")
    for i, col in enumerate(s.columns, 1):
        s.column_dimensions[get_column_letter(i)].width = min(76, max(18,
            max(len(str(c.value or "")) for c in col[:120])+2))
    w.save(OUT/name)


# Rule, category, original evidence, implementation, audit conclusion.
constraints = [
 ("Q1 single service round trip", "SOURCE_EXPLICIT", "DOCX Q1; P046", "src/relief_uav/q1/solver.py:129", "Implemented as one destination per trip"),
 ("Q1 indivisible boxes exactly once", "SOURCE_EXPLICIT", "DOCX Q1", "src/relief_uav/q1/solver.py:85; src/relief_uav/validation/q1.py", "Subset-cover and independent validation"),
 ("Q1 mass and volume capacity", "SOURCE_EXPLICIT", "transport workbook", "src/relief_uav/q1/solver.py:107", "Both checked"),
 ("Q1 safety SOC margin", "SOURCE_EXPLICIT", "P056-P058", "src/relief_uav/q1/solver.py:40", "20% baseline; sensitivity analysis"),
 ("Q1 lexicographic trip/energy/time order", "PROJECT_ASSUMPTION", "Q1 asks multiobjective, no scalar order", "src/relief_uav/q1/solver.py:154", "Document trade-off, not unique official optimum"),
 ("Q2 multi-stop route and shared physical resources", "SOURCE_EXPLICIT", "DOCX Q2; P062", "src/relief_uav/q2/scheduler_v2.py:68-110", "Transport UAV and compatible battery no-overlap"),
 ("Q2 medical desired deadline", "SOURCE_EXPLICIT", "DOCX P018", "src/relief_uav/q2/scheduler_v2.py:100; src/relief_uav/validation/q2.py:135", "Hard constraint"),
 ("Q2 first-batch deadline", "SOURCE_EXPLICIT", "DOCX P018", "src/relief_uav/q2/scheduler_v2.py:102; src/relief_uav/validation/q2.py:136", "Hard constraint"),
 ("Q2 ordinary desired deadline", "SOURCE_EXPLICIT", "DOCX P018", "src/relief_uav/q2/scheduler_v2.py:104", "Soft weighted tardiness"),
 ("Q2 battery available only after full recharge", "SOURCE_EXPLICIT", "P062-P065", "src/relief_uav/q2/scheduler_v2.py:85-93", "Occupied through charge end"),
 ("Q2 epsilon selection", "PROJECT_ASSUMPTION", "No official scalarization", "src/relief_uav/q2/objectives.py:42", "Transparent Pareto representative rule"),
 ("Q3 continuous transport communication", "SOURCE_EXPLICIT", "P070-P101", "src/relief_uav/communication/coverage.py:76; src/relief_uav/validation/q3.py:109", "Numerically checked at 0.25 s and transitions"),
 ("Q3 direct priority and one-hop relay", "SOURCE_EXPLICIT", "P070-P101", "src/relief_uav/communication/coverage.py:61", "Implemented"),
 ("Q3 relay drone preparation/turnaround", "SOURCE_EXPLICIT", "P059-P060", "src/relief_uav/q3/joint_cp.py:245-259", "Occupied through turnaround"),
 ("Q3 component recharge and reuse", "SOURCE_EXPLICIT", "P062-P065", "src/relief_uav/q3/joint_cp.py:203-260; src/relief_uav/validation/q3.py:106", "Experimental CP now assigns one of six components and occupies through full-charge end"),
 ("Q3 one component permanently per relay sortie", "ALGORITHM_RESTRICTION", "Contradicts P062-P065 reuse", "main@92f39ab src/relief_uav/q3/joint_cp.py:86,315", "Previous algorithm artificially capped sorties at six; removed on experiment branch"),
 ("Q3 maximum six relay sorties", "ALGORITHM_RESTRICTION", "Six components are stock, not sortie count", "main@92f39ab src/relief_uav/q3/joint_cp.py:83", "Previous default search cap; configurable 6/8/10/12 in experiment"),
 ("Q3 ordinary desired time forced when seed J1=0", "ALGORITHM_RESTRICTION", "P018 says soft", "main@92f39ab src/relief_uav/q3/joint_cp.py:237", "Zero mode retained; Pareto mode removes this hardening"),
 ("Q3 hover all configured AGL levels", "ALGORITHM_RESTRICTION", "P022 allows DEM interior up to 300 m", "main@92f39ab src/relief_uav/q3/candidates.py:84", "Prior coarse stage searched only maximum; full mode searches every level"),
 ("Q3 hover XY in all valid DEM", "ALGORITHM_RESTRICTION", "P022 entire DEM", "main@92f39ab src/relief_uav/q3/candidates.py:34", "Prior local rectangle; buffered and full DEM modes added"),
 ("Q3 route archive top 25 scalar", "ALGORITHM_RESTRICTION", "No source route count cap", "src/relief_uav/q3/search.py:157", "Experimental multiobjective/diversity mode added"),
 ("Q3 candidate coverage using five samples", "POTENTIAL_BUG", "Continuous coverage required", "src/relief_uav/q3/joint_cp.py:49; src/relief_uav/validation/q3.py:109", "Search approximation only; reject any fine-validator outage"),
 ("Q3 J1/J2 lexicographic proxy", "PROJECT_ASSUMPTION", "P022-P024 list four objectives", "src/relief_uav/q3/joint_cp.py:291", "Experimental epsilon plus energy/sortie objectives"),
 ("Q4 fixed Q3 tasks", "SOURCE_EXPLICIT", "P026-P028", "src/relief_uav/q4/inheritance.py:15; src/relief_uav/validation/q4.py:129", "Frozen projection and SHA"),
 ("Q4 strict relay indivisibility", "PROJECT_ASSUMPTION", "P026-P028 group interpretation ambiguous", "src/relief_uav/q4/inheritance.py:70; src/relief_uav/q4/solver.py:39", "Strict policy; compare alternatives"),
 ("Q4 full relay replication", "PROJECT_ASSUMPTION", "Not uniquely specified", "src/relief_uav/q4/solver.py:38", "Sensitivity policy"),
 ("Q4 exact K2/K3 partition enumeration", "PROJECT_ASSUMPTION", "P026-P028 demand K=2,3, not an algorithm", "src/relief_uav/q4/partition.py; src/relief_uav/q4/solver.py:62", "Exact enumeration is a sound method for current scale"),
 ("Q4 resource gap/scale sum", "PROJECT_ASSUMPTION", "No unique scalarization", "src/relief_uav/q4/objectives.py:8-31", "Retain full candidate set and sensitivity"),
 ("Q4 Gap→Scale→CV selection", "PROJECT_ASSUMPTION", "No official lexicographic order", "src/relief_uav/q4/objectives.py:35", "Can select highly unbalanced representative"),
 ("Q4 workload sum mission durations", "PROJECT_ASSUMPTION", "No unique workload formula", "src/relief_uav/q4/resources.py:122", "Separate multidimensional workload sensitivity"),
]

formulas = [
 ("Horizontal WGS84 distance", "SOURCE_EXPLICIT", "P046", "pyproj.Geod inverse, m", "src/relief_uav/geo/dem.py:35", "PASS"),
 ("DEM maximum over touched pixels", "SOURCE_EXPLICIT", "P046", "supercover all closed-segment pixels, m", "src/relief_uav/geo/dem.py:47", "PASS"),
 ("Cruise MSL", "SOURCE_EXPLICIT", "P046", "max DEM ground + 50 m", "src/relief_uav/geo/segments.py:63", "PASS"),
 ("Node operating MSL", "SOURCE_EXPLICIT", "P046", "O01 attached elevation; service attached elevation +30 m", "src/relief_uav/data/models.py", "PASS"),
 ("Equivalent range", "SOURCE_EXPLICIT", "P048", "L0-(L0-LF)(q/Q)^(3/2), m", "src/relief_uav/physics/flight.py:21", "PASS"),
 ("Flight time", "SOURCE_EXPLICIT", "P051", "climb/v_up + distance/v_cruise + descent/v_down, s", "src/relief_uav/physics/flight.py:32", "PASS"),
 ("Transport total energy", "SOURCE_EXPLICIT", "P054", "E_hor+E_up, kWh", "src/relief_uav/physics/energy.py:33", "Only sum is in source"),
 ("Transport horizontal energy", "PROJECT_ASSUMPTION", "No component formula in P053-P055 OMML", "E_use*d/L(q), kWh", "src/relief_uav/physics/energy.py:59", "Dimensionally valid; source formula unresolved"),
 ("Transport ascent energy", "PROJECT_ASSUMPTION", "No component formula in P053-P055 OMML", "m*g*h/(eta*3.6e6), kWh", "src/relief_uav/physics/energy.py:60", "Dimensionally valid; source formula unresolved"),
 ("Return SOC", "SOURCE_EXPLICIT", "P057-P058", "1-E/E_use, unitless", "src/relief_uav/physics/battery.py:10", "PASS"),
 ("Two-stage charging", "SOURCE_EXPLICIT", "P063-P065", "<0.9: T[.65(.9-s)/.9+.35]; >=.9: T*.35(1-s)/.1, s", "src/relief_uav/physics/battery.py:23", "PASS"),
 ("Relay cruise energy", "SOURCE_EXPLICIT", "P060", "P_cruise*t/3600, kWh", "src/relief_uav/q3/relay_physics.py:38", "PASS"),
 ("Relay climb energy", "PROJECT_ASSUMPTION", "P060 refers to missing same relation", "m*g*h/(eta*3.6e6), kWh", "src/relief_uav/q3/relay_physics.py:39", "Source formula unresolved"),
 ("Relay setup energy", "PROJECT_ASSUMPTION", "P060 does not specify setup power state", "(P_hover+P_comm)*setup/3600", "src/relief_uav/q3/relay_physics.py:56", "Explicit selectable assumption"),
 ("Radio receiver threshold", "SOURCE_EXPLICIT", "P070-P101", "sensitivity + fade margin, dBm", "src/relief_uav/communication/link_budget.py:8", "PASS"),
 ("Bidirectional link limit", "SOURCE_EXPLICIT", "P070-P101", "min of both directions, dB", "src/relief_uav/communication/link_budget.py:19", "PASS"),
 ("FSPL", "SOURCE_EXPLICIT", "P070-P101", "32.45+20log10(f_MHz)+20log10(D_km), dB", "src/relief_uav/communication/link_budget.py:25", "PASS"),
 ("Obstruction", "SOURCE_EXPLICIT", "P070-P101", "add attachment obstruction loss, dB", "src/relief_uav/communication/coverage.py:35", "PASS"),
 ("LOS terrain", "SOURCE_EXPLICIT", "P070-P101", "30 m DEM cell intersection", "src/relief_uav/communication/terrain_los.py:29", "Conservative pixel-plate interpretation"),
]

assumptions = [
 ("Transport energy components", "PROJECT_ASSUMPTION", "Keep pending source clarification; sensitivity, no claim of official formula"),
 ("Relay ascent energy", "PROJECT_ASSUMPTION", "Uses same assumed mass-height-energy relation"),
 ("Relay setup power", "PROJECT_ASSUMPTION", "hover_plus_comm; hover_only sensitivity available"),
 ("Q1 objective priority", "PROJECT_ASSUMPTION", "Trip count then energy/time representative"),
 ("Q2/Q3 Pareto selection", "PROJECT_ASSUMPTION", "Explicit epsilon and lexicographic representative"),
 ("Q3 max relay sortie slots", "ALGORITHM_RESTRICTION", "6 default is search parameter, not physical stock; experiment 6/8/10/12"),
 ("Q3 hover candidate pool", "ALGORITHM_RESTRICTION", "Finite XY/AGL grid approximates continuous domain"),
 ("Q3 communication atom eligibility", "ALGORITHM_RESTRICTION", "Five spatial samples; fine independent validator required"),
 ("Q4 strict relay ownership", "PROJECT_ASSUMPTION", "Ambiguous whether cross-group relay may be copied or trimmed"),
 ("Q4 workload and selection", "PROJECT_ASSUMPTION", "Not prescribed formula; multiple metrics and choices reported"),
]

workbook("constraint_audit.xlsx", ["Rule", "Tag", "Original evidence", "Code location", "Finding"], constraints)
workbook("formula_audit.xlsx", ["Formula", "Tag", "Original evidence", "Formula/unit", "Code location", "Finding"], formulas)
workbook("modeling_assumptions.xlsx", ["Assumption", "Tag", "Treatment"], assumptions)

report = """# Q1–Q4 source-to-code audit, experiment branch

Source precedence: original DOCX and attachments, then TASK.md and repository docs.
The DOCX XML contains 72 `m:oMath` objects and 12 math paragraphs. Its P053–P055
specify only `E_total = E_horizontal + E_up`; no missing component relationship was
recovered. The existing horizontal and climb formulas remain explicit project
assumptions. They are dimensionally consistent but cannot be called official.

The physical resource error in the prior Q3 *optimizer* was equating six energy
components with at most six lifetime relay sorties and assigning a never-reused
component to every sortie. P062–P065 allow reuse after charging to 100%.
The prior independent Q3 validator already checked component intervals through
charge completion, so a validator PASS did not reveal the search restriction.
The experiment branch models component assignment and charge-end NoOverlap.

The zero-tardiness mode is a deliberate lexicographic search restriction. It
becomes incorrect if presented as a source hard deadline for ordinary cargo.
The experiment's Pareto mode keeps only medical and first-batch deadlines hard.
The existing local hover rectangle, maximum-altitude coarse search and scalar
route archive are finite algorithm restrictions, not source constraints.

Q4's strict relay ownership, full replication, workload sum, resource-gap
scalarization and lexicographic representative are documented project choices.
The exact partition enumeration itself is appropriate for the current scale.
Comparison experiments are under `outputs/experiments/q3_q4_improvement/`.
No formal Q1–Q4 result is replaced on this branch.
Quantitative experiments and validator statuses are in
`outputs/experiments/q3_q4_improvement/audit_report.md`.
"""
(OUT/"audit_report.md").write_text(report, encoding="utf-8")
print(OUT)
