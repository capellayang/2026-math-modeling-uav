"""Produce a human-checkable workbook from official nodes and common flight rules."""

import argparse
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from relief_uav.data import load_scenario
from relief_uav.geo import build_segment_matrix
from relief_uav.physics import flight_phases


def build(force_recompute: bool = False) -> Path:
    scenario = load_scenario(ROOT)
    matrix = build_segment_matrix(scenario, force_recompute=force_recompute)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "说明"
    overview.append(["项目", "说明"])
    overview.append(["计算范围", "基础几何与飞行时间；不含能耗、载荷或优化结果"])
    overview.append(["距离", "WGS84椭球两节点测地距离，单位m"])
    overview.append(["航线像元", "EPSG:4326两节点经纬度直线的全部相交像元（含边界接触）"])
    overview.append(["DEM", "GeoTIFF主数据；-32767按附件显式识别为NoData"])
    overview.append(["节点作业高度", "O01采用附件海拔；服务区采用附件海拔+30m"])
    overview.append(["缓存", str(matrix.cache_path)])
    overview.append(["缓存读取", matrix.loaded_from_cache])
    overview.append(["航段数", len(matrix.records)])
    overview.append(["能耗", "E_hor与E_up原题缺式，故本文件不计算架次能耗或SOC"])
    overview.append(["多点示例", "O01→S001→S005→O01；每段重新计算DEM和作业高度"])
    ws = workbook.create_sheet("航段核验")
    ws.append(["路线", "航段序号", "起点", "终点", "水平距离（m）", "最高DEM（m）",
               "巡航海拔（m）", "起点作业海拔（m）", "终点作业海拔（m）",
               "爬升（m）", "下降（m）", "相交像元数", "机型", "爬升时间（s）",
               "巡航时间（s）", "下降时间（s）", "航段总飞行时间（s）"])
    routes = [("单点往返", ["O01", "S001", "O01"]),
              ("多点示例", ["O01", "S001", "S005", "O01"])]
    for route_name, route in routes:
        for index, (start, end) in enumerate(zip(route, route[1:]), 1):
            segment = matrix.get(start, end)
            for model in scenario.transport_models.values():
                phases = flight_phases(model, segment)
                ws.append([route_name, index, start, end, segment.horizontal_distance_m,
                           segment.maximum_dem_m, segment.cruise_altitude_m,
                           segment.start_operating_altitude_m, segment.end_operating_altitude_m,
                           segment.climb_m, segment.descent_m, segment.touched_cell_count,
                           model.model_id, phases.climb_s, phases.cruise_s,
                           phases.descent_s, phases.total_s])
    overview.column_dimensions["A"].width = 22
    overview.column_dimensions["B"].width = 92
    for col, width in enumerate([16, 12, 13, 13, 20, 19, 19, 23, 23, 17, 17, 15,
                                 11, 19, 19, 19, 24], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="24445C")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row[4:11] + row[13:17]:
            cell.number_format = "0.000000"
    out = ROOT / "outputs/validation/base_physics_check.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    print(f"Wrote {out}; directed pairs={len(matrix.records)}; loaded_from_cache={matrix.loaded_from_cache}")
    first = matrix.get("O01", "S001")
    print("O01→S001", first)
    for model in scenario.transport_models.values():
        print(model.model_id, flight_phases(model, first))
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--force-recompute", action="store_true", help="Ignore existing node-pair cache")
    build(parser.parse_args().force_recompute)
