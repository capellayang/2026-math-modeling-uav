"""Audit CURRENT, seed 20260925 and seed 20260927, then freeze formal Q3/Q4.

The script builds and validates every export in an isolated staging directory
before copying it over the formal outputs. It never runs a new optimizer search.
"""

from dataclasses import asdict
from hashlib import sha256
import json
from math import isclose
from pathlib import Path
import shutil
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from relief_uav.communication.coverage import RadioEnvironment
from relief_uav.data import load_scenario
from relief_uav.geo import build_segment_matrix
from relief_uav.geo.dem import DigitalElevationModel
from relief_uav.geo.segments import dem_source_path
from relief_uav.q3.model import Q3AlgorithmConfig
from relief_uav.q3.report import load_q3_solution, save_q3_outputs
from relief_uav.q3.solver import _candidate_row
from relief_uav.q4.model import Q4AlgorithmConfig
from relief_uav.q4.solver import solve_q4
from relief_uav.q4.report import save_q4_outputs, save_q4_validation
from relief_uav.validation.q3 import validate_q3
from relief_uav.validation.q4 import validate_q4


OUT = ROOT / "outputs" / "experiments" / "final_freeze"
STAGE = OUT / "staging_root"
REFERENCE = OUT / "reference_current"
OLD_Q3 = REFERENCE / "q3_summary.json"
OLD_Q4 = REFERENCE / "q4_summary.json"
OLD_Q4_TEMPLATE = REFERENCE / "结果提交_Q4.xlsx"
SEEDS = ROOT / "outputs" / "experiments" / "q3_five_seed"
FIXED_BASE = "12859c21794f1902c2bc990c7f9d75f58750f435"
OLD_FORMAL_COMMIT = "d1144d26d83cd62c0ac013706c7d5c0f32bf7809"


def _sha(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2,
                               allow_nan=False), encoding="utf-8")


def _source_fingerprint() -> str:
    digest = sha256()
    for path in sorted((ROOT / "src" / "relief_uav" / "q3").glob("*.py")):
        digest.update(path.name.encode())
        digest.update(path.read_bytes())
    return digest.hexdigest()


def _restore_reference_current() -> None:
    """Pin CURRENT to the pre-freeze commit, even on later reruns."""
    REFERENCE.mkdir(parents=True, exist_ok=True)
    for relative, target in (
        ("outputs/q3/q3_summary.json", OLD_Q3),
        ("outputs/q4/q4_summary.json", OLD_Q4),
        ("outputs/q4/结果提交_Q4.xlsx", OLD_Q4_TEMPLATE),
    ):
        result = subprocess.run(["git", "show", f"{OLD_FORMAL_COMMIT}:{relative}"],
                                cwd=ROOT, capture_output=True, check=True)
        data = result.stdout
        # These JSON worktree files were checked out with CRLF on Windows;
        # Q4 binds the exact Q3 *file bytes*, not normalized JSON content.
        if target.suffix == ".json" and b"\r\n" not in data:
            data = data.replace(b"\n", b"\r\n")
        target.write_bytes(data)


def _assert_unchanged_model() -> dict:
    workbook = load_workbook(SEEDS / "q3_five_seed_comparison.xlsx",
                             read_only=True, data_only=True)
    fixed = {key: value for key, value in list(
        workbook["Fixed configuration"].values)[1:]}
    assert fixed["q3_source_sha256"] == _source_fingerprint(), "Q3 source changed after seed run"
    baseline = ROOT / "outputs" / "q3" / "baseline_q2_v2_summary.json"
    assert fixed["q2_baseline_sha256"] == _sha(baseline), "Q2 source changed after seed run"
    changed = subprocess.run(["git", "diff", "--name-only", FIXED_BASE, "--",
        "src/relief_uav/q2", "src/relief_uav/q3", "src/relief_uav/q4",
        "src/relief_uav/physics", "src/relief_uav/communication",
        "src/relief_uav/geo", "src/relief_uav/validation"],
        cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
    assert not changed, f"Formula/constraint/validator code changed: {changed}"
    return {"q3_source_sha256": fixed["q3_source_sha256"],
            "q2_baseline_sha256": fixed["q2_baseline_sha256"],
            "model_diff_from_experiment_base": "none"}


def _numbers(q3, validation, q4) -> dict:
    result = {"J1": q3.objective.weighted_tardiness,
        "J2_s": q3.objective.joint_makespan_s,
        "J3_kwh": q3.objective.joint_energy_kwh,
        "transport_energy_kwh": q3.objective.transport_energy_kwh,
        "relay_energy_kwh": q3.objective.relay_energy_kwh,
        "NT": q3.objective.transport_sortie_count,
        "NR": q3.objective.relay_sortie_count,
        "boxes": len(q3.transport.box_deliveries),
        "outage_s": validation.outage_s,
        "direct_s": validation.direct_communication_s,
        "relay_s": validation.relay_communication_s,
        "minimum_relay_access_margin_db": validation.minimum_relay_access_margin_db,
        "minimum_relay_backhaul_margin_db": validation.minimum_relay_backhaul_margin_db,
        "strict_components": len(q4.strict_components),
        "largest_strict_component": max(len(c.service_areas) for c in q4.strict_components),
        "strict_components_detail": [list(c.service_areas) for c in q4.strict_components],
        "setup_energy_mode": q3.setup_energy_mode}
    for p in q4.policy_results:
        chosen = next((c for c in p.candidates if c.selected), None)
        key = ("strict" if p.relay_policy == "strict_no_duplication" else "replicate")
        result[f"{key}_K{p.k}"] = {
            "legal_partitions": len(p.candidates),
            "gap": chosen.stock_gap_total if chosen else None,
            "scale": chosen.resource_scale if chosen else None,
            "cv": chosen.workload_cv if chosen else None,
            "resource_vector": chosen.resource_vector if chosen else None,
            "group_signature": chosen.canonical_group_signature if chosen else None}
    return result


def _same(a, b) -> bool:
    if isinstance(a, (float, int)) and isinstance(b, (float, int)):
        return isclose(a, b, rel_tol=0, abs_tol=1e-7)
    return a == b


def _nonempty_rows(ws):
    return [tuple(row) for row in ws.values if any(value is not None for value in row)]


def _audit_q3_exports(q3, validation) -> dict:
    q3_dir = STAGE / "outputs" / "q3"
    template = load_workbook(ROOT / "结果提交模板.xlsx", read_only=True, data_only=True)
    exported = load_workbook(q3_dir / "结果提交_Q3.xlsx", read_only=True, data_only=True)
    assert exported.sheetnames == template.sheetnames
    for name in template.sheetnames:
        if name not in ("Q3_中继架次", "Q3_通信保障"):
            assert list(exported[name].values) == list(template[name].values), name
    expected_relay = [[r.sortie_id, r.drone_id, r.component_id,
        r.preparation_start_s, r.hover.longitude_deg, r.hover.latitude_deg,
        r.hover.hover_msl_m, r.link_setup_end_s, r.service_end_s,
        r.return_o01_s, r.total_energy_kwh] for r in q3.relays]
    relay_rows = _nonempty_rows(exported["Q3_中继架次"])[1:]
    assert len(relay_rows) == len(expected_relay)
    for actual, expected in zip(relay_rows, expected_relay):
        assert all(_same(a, b) for a, b in zip(actual[:11], expected)), "Q3 relay template mismatch"
    expected_comm = [[r.transport_sortie_id, r.phase, r.start_s, r.end_s,
        "直连" if r.mode == "DIRECT" else "中继", r.relay_sortie_id]
        for r in validation.intervals]
    comm_rows = _nonempty_rows(exported["Q3_通信保障"])[1:]
    assert len(comm_rows) == len(expected_comm)
    for actual, expected in zip(comm_rows, expected_comm):
        assert all(_same(a, b) for a, b in zip(actual[:6], expected)), "Q3 coverage template mismatch"
    counts = {"q3_transport_sorties.xlsx": len(q3.transport.sorties),
        "q3_box_deliveries.xlsx": len(q3.transport.box_deliveries),
        "q3_relay_sorties.xlsx": len(q3.relays),
        "q3_relay_resource_timeline.xlsx": len(q3.relays),
        "q3_communication_coverage.xlsx": len(validation.intervals),
        "q3_pareto.xlsx": q3.pareto_count}
    for filename, expected in counts.items():
        book = load_workbook(q3_dir / filename, read_only=True, data_only=True)
        actual = len(_nonempty_rows(book.worksheets[0])) - 1
        assert actual == expected, f"{filename}: {actual} != {expected}"
    restored = load_q3_solution(q3_dir / "q3_summary.json")
    assert asdict(restored) == asdict(q3), "Q3 summary changes frozen decisions"
    return {"official_relay_rows": len(relay_rows),
            "official_communication_rows": len(comm_rows),
            "supplemental_row_counts": counts,
            "non_Q3_template_sheets_unchanged": True}


def _copy_files(source: Path, destination: Path) -> None:
    for path in source.rglob("*"):
        if path.is_file():
            target = destination / path.relative_to(source)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target)


def _mark_unavailable_history() -> None:
    """Avoid presenting an invented ALNS trajectory as freeze evidence."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(9, 4))
    ax.axis("off")
    ax.text(.5, .5, "Historical ALNS trace was not archived by the five-seed batch.\n"
            "Selected solution: 2000 iterations; see final_freeze report.",
            ha="center", va="center", fontsize=11, transform=ax.transAxes)
    fig.tight_layout()
    fig.savefig(STAGE / "outputs/q3/figures/q3_objective_history.png", dpi=170)
    plt.close(fig)


def _comparison_book(metrics: dict[str, dict], checks: list[tuple]) -> None:
    book = Workbook()
    ws = book.active
    ws.title = "CURRENT vs 20260927"
    ws.append(("metric", "CURRENT", "20260925 control", "20260927 formal"))
    fields = ("J1", "J2_s", "J3_kwh", "transport_energy_kwh",
              "relay_energy_kwh", "NT", "NR", "boxes", "outage_s",
              "direct_s", "relay_s", "minimum_relay_access_margin_db",
              "minimum_relay_backhaul_margin_db", "strict_components",
              "largest_strict_component", "setup_energy_mode")
    for field in fields:
        ws.append((field, *(metrics[name][field]
                  for name in ("CURRENT", "20260925", "20260927"))))
    for policy in ("strict", "replicate"):
        for k in (2, 3):
            for field in ("legal_partitions", "gap", "scale", "cv",
                          "resource_vector", "group_signature"):
                key = f"{policy}_K{k}"
                values = [metrics[name][key][field]
                          for name in ("CURRENT", "20260925", "20260927")]
                ws.append((f"{key}.{field}", *(json.dumps(v, ensure_ascii=False)
                          if isinstance(v, dict) else v for v in values)))
    cs = book.create_sheet("Freeze checks")
    cs.append(("check", "status", "evidence"))
    for item in checks:
        cs.append(item)
    for sheet in book:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24445C")
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                100, max(18, max(len(str(c.value or "")) for c in column) + 2))
    book.save(OUT / "final_freeze_comparison.xlsx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    _restore_reference_current()
    provenance = _assert_unchanged_model()
    scenario = load_scenario(ROOT)
    segments = build_segment_matrix(scenario)
    paths = {"CURRENT": OLD_Q3,
             "20260925": SEEDS / "20260925" / "q3.json",
             "20260927": SEEDS / "20260927" / "q3.json"}
    q3s, validations, q4s, metrics = {}, {}, {}, {}
    checks = [("Q3 source and Q2 baseline fingerprints", "PASS",
               json.dumps(provenance, ensure_ascii=False)),
              ("Q2/Q3/Q4 physics/communication/geo/validators unchanged", "PASS",
               f"git diff {FIXED_BASE[:7]} -- src/... is empty")]
    for name, path in paths.items():
        print(f"FREEZE_Q3_CHECK {name}", flush=True)
        q3 = load_q3_solution(path)
        env = RadioEnvironment(scenario, DigitalElevationModel(dem_source_path(ROOT)))
        check = validate_q3(env, segments, q3, max_step_s=0.25,
                            transition_tolerance_s=0.05)
        assert check.passed and check.outage_s <= 1e-9, (name, check.issues[:10])
        box_ids = [d.box_id for d in q3.transport.box_deliveries]
        assert len(box_ids) == len(set(box_ids)) == len(scenario.boxes) == 80
        assert set(box_ids) == set(scenario.boxes)
        assert q3.setup_energy_mode == "hover_plus_comm"
        q3s[name], validations[name] = q3, check
        checks.append((f"{name} Q2/Q3 source physics, SOC, charge, resources, radio", "PASS",
                       f"80 unique boxes; 0.25 s; OUTAGE={check.outage_s}; issues=0"))
        print(f"FREEZE_Q4_CHECK {name}", flush=True)
        q4 = solve_q4(scenario, q3, path,
                      {"passed": True, "outage_s": check.outage_s},
                      Q4AlgorithmConfig(relay_policy="both"))
        payload = json.loads(json.dumps(asdict(q4), ensure_ascii=False))
        q4_check = validate_q4(scenario, q3, path, payload, q3_passed=True,
            official_template_path=OLD_Q4_TEMPLATE
            if name == "CURRENT" else None)
        assert q4_check.passed, (name, q4_check.issues[:10])
        if name == "CURRENT":
            previous = json.loads(OLD_Q4.read_text(encoding="utf-8"))
            previous_check = validate_q4(scenario, q3, path, previous,
                q3_passed=True,
                official_template_path=OLD_Q4_TEMPLATE)
            assert previous_check.passed, previous_check.issues[:10]
            checks.append(("CURRENT saved Q4 and official template", "PASS",
                           f"{previous_check.checked_candidates} candidates"))
        checks.append((f"{name} Q4 exact enumeration and frozen inheritance", "PASS",
                       f"{q4_check.checked_candidates} candidates; issues=0"))
        q4s[name] = q4
        metrics[name] = _numbers(q3, check, q4)
        print(f"CHECKED {name}: J2={metrics[name]['J2_s']:.3f} "
              f"J3={metrics[name]['J3_kwh']:.6f}", flush=True)
    assert metrics["20260927"]["J1"] == 0
    assert metrics["20260927"]["J2_s"] < metrics["CURRENT"]["J2_s"]
    assert metrics["20260927"]["J3_kwh"] < metrics["CURRENT"]["J3_kwh"]
    # Staging is a complete export with a freshly bound Q4 source fingerprint.
    (STAGE / "outputs/q3").mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / "结果提交模板.xlsx", STAGE / "结果提交模板.xlsx")
    shutil.copy2(ROOT / "outputs/q3/baseline_q2_v2_summary.json",
                 STAGE / "outputs/q3/baseline_q2_v2_summary.json")
    shutil.copy2(ROOT / "outputs/q3/baseline_q2_direct_audit.json",
                 STAGE / "outputs/q3/baseline_q2_direct_audit.json")
    shutil.copy2(ROOT / "outputs/q3/baseline_q2_direct_audit.xlsx",
                 STAGE / "outputs/q3/baseline_q2_direct_audit.xlsx")
    selected = q3s["20260927"]
    fine = validations["20260927"]
    pareto_row = {**_candidate_row("Q3-SEED-20260927", selected, fine),
                  "selected": True}
    direct_audit = json.loads((SEEDS / "20260927/cache/baseline_q2_direct_audit.json")
                              .read_text(encoding="utf-8"))
    save_q3_outputs(STAGE, RadioEnvironment(scenario,
        DigitalElevationModel(dem_source_path(ROOT))), selected, fine,
        Q3AlgorithmConfig(seed=20260927), (pareto_row,), (), direct_audit)
    _mark_unavailable_history()
    q3_export = _audit_q3_exports(selected, fine)
    checks.append(("Q3 official and supplemental export mapping", "PASS",
                   json.dumps(q3_export, ensure_ascii=False)))
    staged_q3 = STAGE / "outputs/q3/q3_summary.json"
    staged_q4 = solve_q4(scenario, selected, staged_q3,
        {k: v for k, v in asdict(fine).items() if k != "intervals"},
        Q4AlgorithmConfig(relay_policy="both"))
    staged_payload = json.loads(json.dumps(asdict(staged_q4), ensure_ascii=False))
    preflight = validate_q4(scenario, selected, staged_q3, staged_payload, q3_passed=True)
    assert preflight.passed, preflight.issues[:10]
    q4_dir = save_q4_outputs(STAGE, scenario, staged_q4, preflight)
    official = validate_q4(scenario, selected, staged_q3, staged_payload,
        q3_passed=True, official_template_path=q4_dir / "结果提交_Q4.xlsx")
    assert official.passed, official.issues[:10]
    save_q4_validation(q4_dir, official)
    checks.append(("Q4 both policies, official template and supplemental exports", "PASS",
                   f"{official.checked_candidates} candidates; {official.checked_groups} groups"))
    _comparison_book(metrics, checks)
    _json(OUT / "freeze_manifest.json", {"status": "STAGED_VALIDATED",
        "provenance": provenance,
        "current_q3_sha256": _sha(OLD_Q3), "current_q4_sha256": _sha(OLD_Q4),
        "staged_q3_sha256": _sha(staged_q3),
        "staged_q4_sha256": _sha(q4_dir / "q4_summary.json"),
        "metrics": metrics, "checks": checks,
        "search_history_note": "Five-seed batch preserved solution and 2000 iteration count, "
            "but not its per-iteration ALNS trajectory. The frozen CSV is header-only; "
            "no historical points were invented. Q3-SEED-20260927 is a provenance "
            "label, not a recovered internal solver candidate ID."})
    print("STAGING_VALIDATED; publishing formal Q3/Q4", flush=True)
    backup = OUT / "previous_formal"
    shutil.copytree(ROOT / "outputs/q3", backup / "q3", dirs_exist_ok=True)
    shutil.copytree(ROOT / "outputs/q4", backup / "q4", dirs_exist_ok=True)
    try:
        _copy_files(STAGE / "outputs/q3", ROOT / "outputs/q3")
        _copy_files(STAGE / "outputs/q4", ROOT / "outputs/q4")
        assert _sha(ROOT / "outputs/q3/q3_summary.json") == _sha(staged_q3)
        assert _sha(ROOT / "outputs/q4/q4_summary.json") == _sha(q4_dir / "q4_summary.json")
        formal_q3 = load_q3_solution(ROOT / "outputs/q3/q3_summary.json")
        assert asdict(formal_q3) == asdict(selected)
        formal_q4 = json.loads((ROOT / "outputs/q4/q4_summary.json").read_text(encoding="utf-8"))
        last = validate_q4(scenario, formal_q3, ROOT / "outputs/q3/q3_summary.json",
            formal_q4, q3_passed=True,
            official_template_path=ROOT / "outputs/q4/结果提交_Q4.xlsx")
        assert last.passed, last.issues[:10]
        assert formal_q4["source_q3_sha256"] == _sha(ROOT / "outputs/q3/q3_summary.json")
    except Exception:
        _copy_files(backup / "q3", ROOT / "outputs/q3")
        _copy_files(backup / "q4", ROOT / "outputs/q4")
        raise
    manifest_path = OUT / "freeze_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["status"] = "FORMAL_Q3_Q4_PUBLISHED_AND_VALIDATED"
    manifest["formal_q3_sha256"] = _sha(ROOT / "outputs/q3/q3_summary.json")
    manifest["formal_q4_sha256"] = _sha(ROOT / "outputs/q4/q4_summary.json")
    manifest["final_q4_checked_candidates"] = last.checked_candidates
    _json(manifest_path, manifest)
    print("FORMAL_FREEZE_PASS", manifest["formal_q3_sha256"], flush=True)


if __name__ == "__main__":
    main()
