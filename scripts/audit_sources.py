"""Read-only source inventory. No optimization or physical model implementation."""
from pathlib import Path
import sys, json, zipfile, hashlib, csv, re
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.audit_deps'))
import openpyxl
from lxml import etree as ET
from pypdf import PdfReader
OUT = ROOT / 'outputs' / 'data_audit'
OUT.mkdir(parents=True, exist_ok=True)
def dump(name, data):
    (OUT/name).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
files = [p for p in list(ROOT.glob('*'))+list((ROOT/'数据').rglob('*')) if p.is_file() and not p.name.startswith('~$') and (p.suffix.lower() in {'.docx','.xlsx','.csv','.mat','.tif','.pdf','.html'} or p.name in {'TASK.md','agent.md','AGENTS.md'})]
dump('source_manifest.json', [{'path':str(p.relative_to(ROOT)), 'bytes':p.stat().st_size, 'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in files])
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def math(n):
    tag=ET.QName(n).localname
    cs=list(n)
    def sub(t):
        e=n.find('m:'+t,ns)
        return math(e) if e is not None else ''
    if tag in ('sSub','sSup','sSubSup'):
        return sub('e')+('_{'+sub('sub')+'}' if tag!='sSup' else '')+('^{'+sub('sup')+'}' if tag!='sSub' else '')
    if tag=='f': return '('+sub('num')+')/('+sub('den')+')'
    if tag=='rad': return 'sqrt('+sub('e')+')'
    if tag=='d':
        pr=n.find('m:dPr',ns)
        def delim(key,default):
            e=pr.find('m:'+key,ns) if pr is not None else None
            return e.get('{'+ns['m']+'}val',default) if e is not None else default
        return delim('begChr','(')+sub('e')+delim('endChr',')')
    if tag=='nary':
        e=n.find('m:naryPr/m:chr',ns)
        return (e.get('{'+ns['m']+'}val') if e is not None else '∫')+'_{'+sub('sub')+'} '+sub('e')
    if tag in ('m','eqArr'): return ' ; '.join(math(c) for c in cs if not ET.QName(c).localname.endswith('Pr'))
    if tag=='t': return n.text or ''
    if tag.endswith('Pr'): return ''
    return ''.join(math(c) for c in cs)
with zipfile.ZipFile(next(ROOT.glob('*.docx'))) as z:
    xml=z.read('word/document.xml'); (OUT/'document.xml').write_bytes(xml)
    tree=ET.fromstring(xml)
    paragraphs=[]; equations=[]
    for i,p in enumerate(tree.findall('.//w:body//w:p',ns),1):
        parts=[]
        for c in p:
            if ET.QName(c).namespace==ns['m']:
                parts.append(' $'+math(c)+'$ ')
            else: parts.extend(c.xpath('.//w:t/text()',namespaces=ns))
        paragraphs.append(f'P{i:03d} '+''.join(parts))
        for m in p.findall('.//m:oMath',ns): equations.append({'paragraph':i,'text':math(m),'xml':ET.tostring(m,encoding='unicode')})
    (OUT/'problem_extracted.txt').write_text('\n'.join(paragraphs),encoding='utf-8')
    dump('office_math.json',equations)
    dump('docx_structure.json',{'math_count':len(equations),'paragraphs':len(paragraphs),'media':[x for x in z.namelist() if x.startswith(('word/media/','word/embeddings/'))], 'objects':len(tree.xpath('//*[local-name()="OLEObject"]'))})
books=[]
for path in sorted(list(ROOT.glob('*.xlsx'))+list((ROOT/'数据').rglob('*.xlsx'))):
    if 'outputs' in path.parts or '.audit_deps' in path.parts or path.name.startswith('~$'): continue
    wb=openpyxl.load_workbook(path,data_only=False)
    sheets=[]
    for ws in wb:
        rows=[[c.value for c in row] for row in ws]
        populated=[i+1 for i,r in enumerate(rows) if any(v is not None for v in r)]
        sheets.append({'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'populated_rows':populated,'merged':[str(r) for r in ws.merged_cells.ranges],'rows':rows,'hidden':ws.sheet_state,'formulas':[{'cell':c.coordinate,'formula':c.value} for row in ws for c in row if c.data_type=='f'],'comments':[{'cell':c.coordinate,'text':c.comment.text} for row in ws for c in row if c.comment]})
    books.append({'file':str(path.relative_to(ROOT)),'sheets':sheets})
dump('workbooks_raw.json',books)
for p in (ROOT/'数据').rglob('*.pdf'):
    reader=PdfReader(p)
    (OUT/'geodata_description.txt').write_text('\n'.join(f'PAGE {i+1}\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages)),encoding='utf-8')
print('Extracted source documents and',len(books),'workbooks')
