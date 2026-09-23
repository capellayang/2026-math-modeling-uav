"""Independent checks of saved Q4 partitions against the immutable Q3 source."""

from dataclasses import asdict, dataclass
import json
from math import isclose
from pathlib import Path

from openpyxl import load_workbook

from relief_uav.data.models import Scenario
from relief_uav.q3.model import Q3Solution
from relief_uav.q4.inheritance import (atomic_components, frozen_q3_projection,
                                        relay_bindings, source_sha256)
from relief_uav.q4.model import RESOURCE_KEYS
from relief_uav.q4.objectives import partition_metrics
from relief_uav.q4.partition import canonical_signature, enumerate_partitions
from relief_uav.q4.resources import global_minimum, stock_vector, task_intervals
from relief_uav.q4.solver import build_task_groups


@dataclass(frozen=True)
class Q4Validation:
    passed: bool
    issues: tuple[str, ...]
    checked_candidates: int
    checked_groups: int
    checked_assignments: int
    source_q3_sha256: str


def _plain(value):
    return json.loads(json.dumps(value, ensure_ascii=False))


def _independent_peak(rows):
    events = sorted((t, order, r.task_id) for r in rows
                    for t, order in ((r.start_s, 1), (r.end_s, 0)))
    active = set()
    maximum = 0
    for _, order, task_id in events:
        if order == 0:
            active.remove(task_id)
        else:
            active.add(task_id)
        maximum = max(maximum, len(active))
    return maximum


def _check_coloring(group, source_group, scenario, q3, issues):
    gid = group["group"]["group_id"]
    expected = task_intervals(scenario, q3, source_group)
    assignments = group["assignments"]
    if len(assignments) != sum(len(v) for v in expected.values()):
        issues.append(f"{gid}: assignment count disagrees with fixed tasks")
    across = {}
    for key in RESOURCE_KEYS:
        rows = expected[key]
        related = [a for a in assignments if a["resource_type"] == key]
        actual = sorted((a["task_id"], a["source_task_id"], a["occupied_start_s"],
                         a["occupied_end_s"], a["source_q3_resource_id"]) for a in related)
        expected_rows = sorted((r.task_id, r.source_task_id, r.start_s, r.end_s,
                                r.source_q3_resource_id) for r in rows)
        if actual != expected_rows:
            issues.append(f"{gid}/{key}: fixed task interval or source ID changed")
        minimum = _independent_peak(rows)
        if group["minimum_recolored_requirement"][key] != minimum:
            issues.append(f"{gid}/{key}: requirement differs from independent peak")
        peaks = [p for p in group["resource_peaks"] if p["resource_type"] == key]
        if len(peaks) != 1 or peaks[0]["required"] != minimum:
            issues.append(f"{gid}/{key}: reported peak wrong")
        by_resource = {}
        for a in related:
            rid = a["internal_resource_id"]
            if not rid.startswith(gid + "-"):
                issues.append(f"{gid}/{key}: internal resource ID crosses group")
            if rid in across and across[rid] != key:
                issues.append(f"{gid}: resource ID reused across different types")
            across[rid] = key
            by_resource.setdefault(rid, []).append(a)
        if len(by_resource) != minimum:
            issues.append(f"{gid}/{key}: coloring uses nonminimum count")
        for rid, uses in by_resource.items():
            ordered = sorted(uses, key=lambda a: a["occupied_start_s"])
            for left, right in zip(ordered, ordered[1:]):
                if left["occupied_end_s"] > right["occupied_start_s"] + 1e-9:
                    issues.append(f"{gid}/{key}/{rid}: overlapping tasks")
    return len(assignments)


def _check_template(path: Path, payload: dict, issues: list[str]):
    if not path.is_file():
        issues.append("Official Q4 workbook missing")
        return
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["Q4_分区配置"]
    expected = []
    for result in payload["policy_results"]:
        if result["relay_policy"] != "strict_no_duplication" or not result["feasible"]:
            continue
        selected = next(c for c in result["candidates"] if c["selected"])
        for g in selected["groups"]:
            r = g["minimum_recolored_requirement"]
            expected.append((result["k"], g["group"]["group_id"],
                ",".join(g["group"]["service_areas"]),
                *(r[key] for key in RESOURCE_KEYS)))
    actual = [tuple(row[:11]) for row in sheet.iter_rows(min_row=2, values_only=True)
              if row[0] is not None]
    if actual != expected:
        issues.append("Official Q4 worksheet differs from selected strict candidates")
    original = path.parents[2] / "结果提交模板.xlsx"
    if original.is_file():
        source = load_workbook(original, read_only=True, data_only=True)
        for name in source.sheetnames:
            if name == "Q4_分区配置":
                continue
            if name not in book.sheetnames or list(source[name].values) != list(book[name].values):
                issues.append(f"Official template non-Q4 worksheet changed: {name}")


def validate_q4(scenario: Scenario, q3: Q3Solution, source_path: Path,
                payload: dict, *, q3_passed: bool,
                official_template_path: Path | None = None) -> Q4Validation:
    issues = []
    fingerprint = source_sha256(source_path)
    if not q3_passed or not payload.get("q3_validation", {}).get("passed"):
        issues.append("Q3 source validator has not passed")
    if payload.get("q3_validation", {}).get("outage_s", 1) > 1e-9:
        issues.append("Q3 source reports communication outage")
    if payload.get("source_q3_sha256") != fingerprint:
        issues.append("Q3 source SHA256 changed")
    if payload.get("source_q3_frozen") != _plain(frozen_q3_projection(q3)):
        issues.append("A frozen Q3 route, box, time, relay, or communication field changed")
    transport_atoms = atomic_components(scenario, q3, strict=False)
    strict_atoms = atomic_components(scenario, q3, strict=True)
    if payload.get("transport_components") != _plain([asdict(c) for c in transport_atoms]):
        issues.append("Transport atomic components differ from source")
    if payload.get("strict_components") != _plain([asdict(c) for c in strict_atoms]):
        issues.append("Strict atomic components differ from source")
    if payload.get("relay_bindings") != _plain(relay_bindings(q3)):
        issues.append("Relay-to-transport bindings differ from source")
    stock = stock_vector(scenario)
    baseline = global_minimum(scenario, q3)
    if payload.get("stock") != stock or payload.get("global_q3_minimum") != baseline:
        issues.append("Official stock or global fixed-Q3 minimum wrong")
    result_by_key = {(r["relay_policy"], r["k"]): r
                     for r in payload.get("policy_results", [])}
    if len(result_by_key) != len(payload.get("policy_results", [])):
        issues.append("Duplicate policy/K result")
    if ("strict_no_duplication", 2) not in result_by_key or (
            "strict_no_duplication", 3) not in result_by_key:
        issues.append("Official strict K2/K3 result missing")
    checked_candidates = checked_groups = checked_assignments = 0
    all_internal_ids = {}
    for (policy, k), result in result_by_key.items():
        atoms = strict_atoms if policy == "strict_no_duplication" else transport_atoms
        expected_partitions = tuple(enumerate_partitions(atoms, k))
        if result["atomic_component_count"] != len(atoms):
            issues.append(f"{policy}/K{k}: atomic component count wrong")
        if not expected_partitions:
            if result["feasible"] or result["candidates"]:
                issues.append(f"{policy}/K{k}: infeasible partition claimed feasible")
            continue
        if not result["feasible"] or len(result["candidates"]) != len(expected_partitions):
            issues.append(f"{policy}/K{k}: incomplete exact enumeration")
        expected_signatures = {canonical_signature(p) for p in expected_partitions}
        found = [c["canonical_group_signature"] for c in result["candidates"]]
        if set(found) != expected_signatures or len(found) != len(set(found)):
            issues.append(f"{policy}/K{k}: missing or duplicate canonical partition")
        expected_ids = [f"{policy}-K{k}-P{i:03d}"
                        for i in range(1, len(expected_partitions) + 1)]
        if [c["candidate_id"] for c in result["candidates"]] != expected_ids:
            issues.append(f"{policy}/K{k}: candidate IDs/order are not canonical")
        for candidate in result["candidates"]:
            checked_candidates += 1
            if candidate["k"] != k or candidate["relay_policy"] != policy:
                issues.append(f"{candidate['candidate_id']}: candidate K/policy mismatch")
            groups = candidate["groups"]
            services = [sid for g in groups for sid in g["group"]["service_areas"]]
            if len(groups) != k or any(not g["group"]["service_areas"] for g in groups):
                issues.append(f"{candidate['candidate_id']}: empty or wrong group count")
                continue
            if sorted(services) != sorted(scenario.services):
                issues.append(f"{candidate['candidate_id']}: service coverage/uniqueness wrong")
                continue
            sites = tuple(tuple(g["group"]["service_areas"]) for g in groups)
            if candidate["canonical_group_signature"] != canonical_signature(sites):
                issues.append(f"{candidate['candidate_id']}: signature wrong")
            try:
                source_groups = build_task_groups(q3, sites, k, policy)
            except (ValueError, KeyError) as exc:
                issues.append(f"{candidate['candidate_id']}: {exc}")
                continue
            for group, source_group in zip(groups, source_groups):
                checked_groups += 1
                if group["group"] != _plain(asdict(source_group)):
                    issues.append(f"{candidate['candidate_id']}/{source_group.group_id}: Q3 task/copy changed")
                checked_assignments += _check_coloring(group, source_group, scenario, q3, issues)
                for a in group["assignments"]:
                    rid = a["internal_resource_id"]
                    owner = all_internal_ids.setdefault((candidate["candidate_id"], rid), source_group.group_id)
                    if owner != source_group.group_id:
                        issues.append(f"{candidate['candidate_id']}: resource ID crosses groups")
                expected_boxes = [b for b in scenario.boxes.values()
                                  if b.service_id in source_group.service_areas]
                if group["box_count"] != len(expected_boxes) or not isclose(
                    group["cargo_mass_kg"], sum(b.mass_kg for b in expected_boxes), abs_tol=1e-8):
                    issues.append(f"{candidate['candidate_id']}: box count or mass changed")
            from relief_uav.q4.resources import group_metrics
            expected_metrics = tuple(group_metrics(scenario, q3, g) for g in source_groups)
            for group, expected_group in zip(groups, expected_metrics):
                comparison = _plain(asdict(expected_group))
                for field in ("service_area_count", "box_count", "cargo_mass_kg",
                              "transport_sortie_count", "relay_task_count",
                              "transport_workload_s", "relay_workload_s",
                              "total_workload_s", "inherited_q3_id_count",
                              "minimum_recolored_requirement", "resource_peaks"):
                    if group[field] != comparison[field]:
                        issues.append(f"{candidate['candidate_id']}/{group['group']['group_id']}: {field} wrong")
            objective = partition_metrics(expected_metrics, stock, baseline)
            for field, value in objective.items():
                if candidate[field] != value:
                    issues.append(f"{candidate['candidate_id']}: {field} wrong")
            if candidate["effective_relay_task_count"] != sum(
                    len(g.relay_tasks) for g in source_groups):
                issues.append(f"{candidate['candidate_id']}: relay copy count wrong")
            if candidate["original_relay_task_count"] != len(q3.relays):
                issues.append(f"{candidate['candidate_id']}: original relay count wrong")
        candidates = result["candidates"]
        if candidates:
            best = min(candidates, key=lambda c: (c["stock_gap_total"],
                c["resource_scale"], c["workload_cv"], c["canonical_group_signature"]))
            marked = [c["candidate_id"] for c in candidates if c["selected"]]
            if marked != [best["candidate_id"]] or result["selected_candidate_id"] != best["candidate_id"]:
                issues.append(f"{policy}/K{k}: lexicographic selection wrong")
            def dominates(a, b):
                va = (a["stock_gap_total"], a["resource_scale"], a["workload_cv"])
                vb = (b["stock_gap_total"], b["resource_scale"], b["workload_cv"])
                return all(x <= y + 1e-10 for x, y in zip(va, vb)) and any(
                    x < y - 1e-10 for x, y in zip(va, vb))
            front = {c["candidate_id"] for c in candidates if not any(
                dominates(other, c) for other in candidates if other is not c)}
            if front != set(result["pareto_candidate_ids"]) or front != {
                    c["candidate_id"] for c in candidates if c["pareto"]}:
                issues.append(f"{policy}/K{k}: Pareto front wrong")
    if official_template_path is not None:
        _check_template(official_template_path, payload, issues)
    return Q4Validation(not issues, tuple(issues), checked_candidates,
                        checked_groups, checked_assignments, fingerprint)
