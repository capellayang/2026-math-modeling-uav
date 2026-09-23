"""Read-only direct-link audit of the committed Q2-v2 transport seed."""

from dataclasses import asdict
from hashlib import sha256
import json
from pathlib import Path

from relief_uav.communication.coverage import RadioEnvironment, coverage_intervals
from relief_uav.q2.report import _book, load_q2_solution


def audit_fingerprint(env: RadioEnvironment, transport) -> str:
    stat = env.dem.path.stat()
    payload = {"dem": (stat.st_size, stat.st_mtime_ns),
               "radio": asdict(env.params),
               "sorties": [(asdict(s.spec), s.preparation_start_s)
                           for s in transport.sorties]}
    return sha256(json.dumps(payload, sort_keys=True,
                             ensure_ascii=False).encode()).hexdigest()


def direct_audit(env: RadioEnvironment, transport, *, step_s: float = 2.0,
                 source: str = "transport candidate"):
    all_intervals = []
    by_sortie = []
    for sortie in transport.sorties:
        rows, _ = coverage_intervals(env, sortie, max_step_s=step_s)
        all_intervals.extend(rows)
        total = sortie.return_o01_time_s-sortie.takeoff_time_s
        direct = sum(r.end_s-r.start_s for r in rows if r.mode == "DIRECT")
        outages = [r for r in rows if r.mode == "OUTAGE"]
        by_sortie.append({"sortie_id": sortie.spec.sortie_id,
                          "total_required_s": total, "direct_s": direct,
                          "direct_ratio": direct/total,
                          "blackout_s": sum(r.end_s-r.start_s for r in outages),
                          "blackout_interval_count": len(outages),
                          "minimum_direct_margin_db": min(r.minimum_margin_db for r in rows)})
    total = sum(x["total_required_s"] for x in by_sortie)
    direct = sum(x["direct_s"] for x in by_sortie)
    result = {"source": source,
              "source_fingerprint": audit_fingerprint(env, transport),
              "verification_max_step_s": step_s,
              "total_required_communication_s": total,
              "direct_communication_s": direct,
              "direct_ratio": direct/total,
              "blackout_s": total-direct,
              "blackout_interval_count": sum(x["blackout_interval_count"] for x in by_sortie),
              "minimum_direct_margin_db": min(x["minimum_direct_margin_db"] for x in by_sortie),
              "sorties": by_sortie,
              "intervals": [asdict(r) for r in all_intervals if r.mode == "OUTAGE"]}
    return result


def save_direct_audit(root: Path, audit: dict) -> None:
    out = root / "outputs/q3"
    out.mkdir(parents=True, exist_ok=True)
    (out / "baseline_q2_direct_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    _book(out / "baseline_q2_direct_audit.xlsx", "逐架次直连覆盖",
          ["运输架次", "需通信（s）", "直连（s）", "直连比例", "失联（s）",
           "失联区间数", "最低直连裕量（dB）"],
          [[r["sortie_id"], r["total_required_s"], r["direct_s"],
            r["direct_ratio"], r["blackout_s"], r["blackout_interval_count"],
            r["minimum_direct_margin_db"]] for r in audit["sorties"]])
    from openpyxl import load_workbook
    workbook = load_workbook(out / "baseline_q2_direct_audit.xlsx")
    sheet = workbook.create_sheet("失联区间")
    sheet.append(["运输架次", "阶段", "开始（s）", "结束（s）", "时长（s）",
                  "最低直连裕量（dB）", "地形遮挡", "失联原因"])
    for r in audit["intervals"]:
        sheet.append([r["transport_sortie_id"], r["phase"], r["start_s"],
                      r["end_s"], r["end_s"]-r["start_s"],
                      r["minimum_margin_db"], r["terrain_blocked"], r["reason"]])
    sheet.freeze_panes = "A2"
    workbook.save(out / "baseline_q2_direct_audit.xlsx")
