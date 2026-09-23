"""Q3 snapshots, auditable workbooks, official-template mapping and figures."""

from collections import Counter
from dataclasses import asdict
import csv
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook

from relief_uav.communication.coverage import RadioEnvironment
from relief_uav.communication.model import CommunicationEndpoint, CommunicationInterval
from relief_uav.communication.trajectory import transport_phases
from relief_uav.q2.report import _book, load_q2_payload, load_q2_solution
from relief_uav.validation.q3 import Q3Validation
from .model import (Q3AlgorithmConfig, Q3Objective, Q3Solution,
                    RelayHoverPoint, RelaySortie, RelayTravel)


def load_q3_solution(path: Path) -> Q3Solution:
    raw = json.loads(path.read_text(encoding="utf-8"))
    relays = []
    for record in raw["relays"]:
        relays.append(RelaySortie(**{**record,
            "hover": RelayHoverPoint(**record["hover"]),
            "travel": RelayTravel(**record["travel"])}))
    return Q3Solution(load_q2_payload(raw["transport"]), tuple(relays),
        tuple(CommunicationInterval(**r) for r in raw["communication"]),
        Q3Objective(**raw["objective"]), raw["seed"], raw["search_seconds"],
        raw["cp_sat_status"], raw["pareto_count"], raw["alns_iterations"],
        raw["setup_energy_mode"])


def _official_template(root: Path, solution: Q3Solution, intervals) -> None:
    target = root / "outputs/q3/结果提交_Q3.xlsx"
    shutil.copy2(root / "结果提交模板.xlsx", target)
    workbook = load_workbook(target)
    relay_sheet = workbook["Q3_中继架次"]
    coverage_sheet = workbook["Q3_通信保障"]
    if [relay_sheet.cell(1, col).value for col in range(1, 12)] != [
        "中继架次编号", "中继无人机编号", "能源组件编号", "开始时刻（s）",
        "悬停经度（°）", "悬停纬度（°）", "悬停海拔（m）", "建链完成时刻（s）",
        "服务结束时刻（s）", "返回O01时刻（s）", "架次能耗（kWh）"]:
        raise ValueError("Official Q3 relay template headers changed")
    if [coverage_sheet.cell(1, col).value for col in range(1, 7)] != [
        "运输架次编号", "通信阶段", "开始时刻（s）", "结束时刻（s）",
        "保障方式", "中继架次编号"]:
        raise ValueError("Official Q3 coverage template headers changed")
    from copy import copy
    def write(ws, index, values):
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col)
            if index > 2 and ws.cell(2, col).has_style:
                cell._style = copy(ws.cell(2, col)._style)
            cell.value = value
    for index, r in enumerate(solution.relays, 2):
        write(relay_sheet, index, [r.sortie_id, r.drone_id, r.component_id,
            r.preparation_start_s, r.hover.longitude_deg, r.hover.latitude_deg,
            r.hover.hover_msl_m, r.link_setup_end_s, r.service_end_s,
            r.return_o01_s, r.total_energy_kwh])
    for index, row in enumerate(intervals, 2):
        if row.mode == "OUTAGE":
            raise ValueError("Official Q3 template cannot contain OUTAGE")
        write(coverage_sheet, index, [row.transport_sortie_id, row.phase,
            row.start_s, row.end_s, "直连" if row.mode == "DIRECT" else "中继",
            row.relay_sortie_id])
    workbook.save(target)


def _plots(root: Path, env: RadioEnvironment, solution: Q3Solution,
           validation: Q3Validation, pareto: tuple[dict, ...], history: tuple[dict, ...]):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    out = root / "outputs/q3/figures"
    out.mkdir(parents=True, exist_ok=True)
    scenario = env.scenario
    fig, ax = plt.subplots(figsize=(10, 8))
    dem = env.dem
    ax.imshow(dem.elevations_m[::8, ::8], extent=(dem.bounds.left, dem.bounds.right,
        dem.bounds.bottom, dem.bounds.top), origin="upper", cmap="terrain", alpha=.55)
    for s in solution.transport.sorties:
        xs = [scenario.nodes[n].longitude_deg for n in s.spec.route]
        ys = [scenario.nodes[n].latitude_deg for n in s.spec.route]
        ax.plot(xs, ys, color="#35608d", alpha=.25, lw=.8)
    ax.scatter([n.longitude_deg for n in scenario.services.values()],
               [n.latitude_deg for n in scenario.services.values()],
               c="#b3261e", s=22, label="Service areas")
    ax.scatter([scenario.dispatch.longitude_deg], [scenario.dispatch.latitude_deg],
               c="#101828", marker="*", s=140, label="O01 / G01")
    ax.scatter([r.hover.longitude_deg for r in solution.relays],
               [r.hover.latitude_deg for r in solution.relays],
               c="#ffd100", marker="^", edgecolor="black", s=95, label="Relay hover")
    ax.set_xlim(109.15,109.31); ax.set_ylim(22.99,23.10)
    ax.set_xlabel("Longitude"); ax.set_ylabel("Latitude")
    ax.legend(); fig.tight_layout(); fig.savefig(out/"q3_transport_relay_map.png",dpi=170);plt.close(fig)
    fig, ax = plt.subplots(figsize=(13, 8))
    identifiers = sorted({r.transport_sortie_id for r in validation.intervals})
    for row in validation.intervals:
        ax.barh(identifiers.index(row.transport_sortie_id), row.end_s-row.start_s,
                left=row.start_s, height=.7,
                color={"DIRECT":"#277da1","RELAY":"#f9c74f","OUTAGE":"#d62828"}[row.mode])
    ax.set_yticks(range(len(identifiers)),identifiers);ax.set_xlabel("Mission time (s)")
    ax.set_title("Transport communication mode");fig.tight_layout()
    fig.savefig(out/"q3_communication_timeline.png",dpi=170);plt.close(fig)
    fig, ax = plt.subplots(figsize=(12, 4))
    for r in solution.relays:
        ax.barh(r.drone_id, r.return_o01_s-r.preparation_start_s,
                left=r.preparation_start_s,color="#577590",alpha=.5)
        ax.barh(r.drone_id,r.service_end_s-r.service_start_s,
                left=r.service_start_s,color="#f8961e",height=.5)
    ax.set_xlabel("Mission time (s)");ax.set_title("Relay flight and service")
    fig.tight_layout();fig.savefig(out/"q3_relay_gantt.png",dpi=170);plt.close(fig)
    fig, ax = plt.subplots(figsize=(12, 4))
    for mode,color in (("DIRECT","#277da1"),("RELAY","#f8961e")):
        rows=[r for r in validation.intervals if r.mode==mode]
        ax.scatter([(r.start_s+r.end_s)/2 for r in rows],
                   [r.minimum_margin_db for r in rows],s=12,color=color,label=mode)
    ax.axhline(0,color="black",lw=.8);ax.set_ylabel("Link margin (dB)")
    ax.set_xlabel("Mission time (s)");ax.legend();fig.tight_layout()
    fig.savefig(out/"q3_link_margin.png",dpi=170);plt.close(fig)
    fig, ax = plt.subplots(figsize=(7,5))
    ax.scatter([r["j1"] for r in pareto],[r["j2_joint_s"] for r in pareto],s=65)
    ax.set_xlabel("J1 weighted tardiness");ax.set_ylabel("J2 joint makespan (s)")
    ax.set_title("Validated Q3 candidates");fig.tight_layout()
    fig.savefig(out/"q3_pareto.png",dpi=170);plt.close(fig)
    fig, ax = plt.subplots(figsize=(9,4))
    ax.plot([r["iteration"] for r in history],[r["score"] for r in history],
            color="#277da1")
    ax.set_xlabel("ALNS iteration");ax.set_ylabel("Best approximate score")
    fig.tight_layout();fig.savefig(out/"q3_objective_history.png",dpi=170);plt.close(fig)


def save_q3_outputs(root: Path, env: RadioEnvironment, solution: Q3Solution,
                    validation: Q3Validation, config: Q3AlgorithmConfig,
                    pareto: tuple[dict, ...], history: tuple[dict, ...],
                    baseline_direct_audit: dict) -> None:
    if not validation.passed:
        raise ValueError("Q3 outputs require independent validator PASS")
    out = root / "outputs/q3"
    out.mkdir(parents=True, exist_ok=True)
    _book(out/"q3_transport_sorties.xlsx", "运输架次",
          ["架次编号","机型","无人机","电池","路线","逐站货箱","准备开始（s）",
           "起飞（s）","返回O01（s）","能耗（kWh）","返航SOC"],
          [[s.spec.sortie_id,s.spec.model_id,s.drone_id,s.battery_id,
            "→".join(s.spec.route),"; ".join(f"{x.service_id}:{','.join(x.box_ids)}"
            for x in s.spec.deliveries),s.preparation_start_s,s.takeoff_time_s,
            s.return_o01_time_s,s.energy_kwh,s.return_soc] for s in solution.transport.sorties])
    _book(out/"q3_box_deliveries.xlsx", "逐箱交付",
          ["货箱编号","架次编号","服务区","交付完成（s）","期望（s）",
           "首批截止（s）","优先系数","加权逾期"],
          [[d.box_id,d.sortie_id,d.service_id,d.delivery_complete_time_s,
            d.desired_delivery_s,d.first_batch_deadline_s,d.emergency_priority,
            d.weighted_tardiness] for d in solution.transport.box_deliveries])
    _book(out/"q3_relay_sorties.xlsx", "中继架次",
          ["架次","中继无人机","能源组件（内部编号）","准备开始（s）","起飞（s）",
           "爬升结束（s）","巡航结束（s）","到达悬停点（s）","建链开始（s）",
           "建链完成（s）","服务开始（s）","服务结束（s）","返程爬升结束（s）",
           "返程巡航结束（s）","返回O01（s）","周转结束（s）","经度（°）",
           "纬度（°）","DEM高程（m）","离地高度（m）","悬停海拔MSL（m）",
           "出航能耗（kWh）","建链能耗（kWh）","服务能耗（kWh）",
           "返航能耗（kWh）","总能耗（kWh）","返航SOC"],
          [[r.sortie_id,r.drone_id,r.component_id,r.preparation_start_s,r.takeoff_s,
            r.outbound_climb_end_s,r.outbound_cruise_end_s,r.arrival_hover_s,
            r.link_setup_start_s,r.link_setup_end_s,r.service_start_s,r.service_end_s,
            r.return_climb_end_s,r.return_cruise_end_s,r.return_o01_s,
            r.turnaround_end_s,r.hover.longitude_deg,r.hover.latitude_deg,
            r.hover.ground_elevation_m,r.hover.hover_agl_m,r.hover.hover_msl_m,
            r.outbound_energy_kwh,r.setup_energy_kwh,r.service_energy_kwh,
            r.return_energy_kwh,r.total_energy_kwh,r.return_soc] for r in solution.relays])
    _book(out/"q3_relay_resource_timeline.xlsx", "实体与组件",
          ["架次","实体","组件","占用开始（s）","返回（s）","实体周转完成（s）",
           "组件开始SOC","组件返回SOC","充电开始（s）","充满（s）"],
          [[r.sortie_id,r.drone_id,r.component_id,r.preparation_start_s,r.return_o01_s,
            r.turnaround_end_s,1.0,r.return_soc,r.charge_start_s,r.charge_end_s]
           for r in solution.relays])
    _book(out/"q3_communication_coverage.xlsx", "通信保障区间",
          ["运输架次","阶段","开始（s）","结束（s）","时长（s）","方式",
           "中继架次","最低裕量（dB）","地形遮挡","直连失联原因"],
          [[r.transport_sortie_id,r.phase,r.start_s,r.end_s,r.end_s-r.start_s,
            r.mode,r.relay_sortie_id,r.minimum_margin_db,r.terrain_blocked,r.reason]
           for r in validation.intervals])
    book = load_workbook(out/"q3_communication_coverage.xlsx")
    sheet=book.create_sheet("链路预算明细")
    sheet.append(["时刻（s）","运输架次","端点A","端点B","三维距离（m）",
                  "地形遮挡","FSPL（dB）","遮挡损耗（dB）","总损耗（dB）",
                  "双向Lmax（dB）","裕量（dB）","可用"])
    by_transport={s.spec.sortie_id:s for s in solution.transport.sorties}
    by_relay={r.sortie_id:r for r in solution.relays}
    phases={sid:transport_phases(env.scenario,s) for sid,s in by_transport.items()}
    for row in validation.intervals:
        t=(row.start_s+row.end_s)/2
        phase=next(p for p in phases[row.transport_sortie_id]
                   if p.name==row.phase and p.start_s-1e-7<=t<=p.end_s+1e-7)
        pos=phase.at(t)
        evaluations=[("Transport","G01",env.link(pos,env.gateway,"direct"))]
        if row.mode=="RELAY":
            relay=by_relay[row.relay_sortie_id]
            hover=CommunicationEndpoint(relay.hover.longitude_deg,
                relay.hover.latitude_deg,relay.hover.hover_msl_m)
            evaluations += [("Transport",relay.sortie_id,env.link(pos,hover,"access")),
                            (relay.sortie_id,"G01",env.link(hover,env.gateway,"backhaul"))]
        for start,end,link in evaluations:
            sheet.append([t,row.transport_sortie_id,start,end,link.distance_3d_m,
                link.terrain_blocked,link.fspl_db,link.obstruction_loss_db,
                link.path_loss_db,link.bidirectional_limit_db,link.margin_db,
                link.available])
    sheet.freeze_panes="A2";book.save(out/"q3_communication_coverage.xlsx")
    _book(out/"q3_pareto.xlsx", "验证候选",
          ["候选","J1","J1归一化","J2联合（s）","运输能耗（kWh）",
           "中继能耗（kWh）","联合能耗（kWh）","运输架次","中继架次",
           "联合架次","直连（s）","中继（s）","中断（s）","CP-SAT","验证","选中"],
          [[p.get(k) for k in ("candidate_id","j1","j1_norm","j2_joint_s",
            "transport_energy_kwh","relay_energy_kwh","joint_energy_kwh",
            "transport_sorties","relay_sorties","joint_sorties","direct_s",
            "relay_s","outage_s","cp_sat_status","validator_status","selected")]
           for p in pareto])
    base=baseline_direct_audit
    baseline=load_q2_solution(out/"baseline_q2_v2_summary.json")
    b=baseline.objective
    o=solution.objective
    _book(out/"q3_algorithm_comparison.xlsx", "基准对比",
          ["方案","J1","运输最晚返航（s）","联合最晚返航（s）","运输能耗（kWh）",
           "中继能耗（kWh）","联合能耗（kWh）","运输架次","中继架次",
           "直连比例","中继比例","中断（s）"],
          [["Q2-v2基准",b.weighted_tardiness,b.makespan_s,b.makespan_s,
            b.total_energy_kwh,0,b.total_energy_kwh,b.sortie_count,0,
            base["direct_ratio"],0,base["blackout_s"]],
           ["Q3运输部分",o.weighted_tardiness,o.transport_makespan_s,
            o.transport_makespan_s,o.transport_energy_kwh,0,o.transport_energy_kwh,
            o.transport_sortie_count,0,None,None,None],
           ["Q3联合",o.weighted_tardiness,o.transport_makespan_s,o.joint_makespan_s,
            o.transport_energy_kwh,o.relay_energy_kwh,o.joint_energy_kwh,
            o.transport_sortie_count,o.relay_sortie_count,
            validation.direct_communication_s/validation.total_required_communication_s,
            validation.relay_communication_s/validation.total_required_communication_s,
            validation.outage_s]])
    summary=asdict(solution)
    summary["validation"]={k:v for k,v in asdict(validation).items() if k!="intervals"}
    summary["config"]=asdict(config)
    summary["pareto_candidates"]=pareto
    summary["selected_candidate_id"]=next(p["candidate_id"] for p in pareto if p["selected"])
    summary["geometry_convention"]=("PROJECT EXTENSION: relay flight MSL = max(DEM max + 50 m, "
        "hover MSL, O01 source elevation); this is not an explicit source formula")
    summary["setup_energy_convention"]=("PROJECT CHOICE: hover_plus_comm during link setup "
        "by default; hover_only is a sensitivity alternative")
    summary["climb_energy_convention"]=("PROJECT MODELING ASSUMPTION: m*g*climb/eta, "
        "not an explicit source energy formula")
    summary["component_id_convention"]="R-COMP-xx are deterministic project-internal IDs"
    alternate = "hover_only" if solution.setup_energy_mode == "hover_plus_comm" else "hover_plus_comm"
    difference_per_sortie = next(iter(env.scenario.relay_models.values())).communication_power_kw * (
        next(iter(env.scenario.relay_models.values())).link_setup_s/3600)
    alternative_energy = (o.joint_energy_kwh - len(solution.relays)*difference_per_sortie
                          if alternate == "hover_only" else
                          o.joint_energy_kwh + len(solution.relays)*difference_per_sortie)
    summary["setup_energy_sensitivity"] = {
        "alternative_mode": alternate,
        "same_schedule_joint_energy_kwh": alternative_energy,
        "same_schedule_soc_feasible": all(
            (r.return_soc + difference_per_sortie/next(iter(env.scenario.relay_models.values())).usable_energy_kwh
             if alternate == "hover_only" else
             r.return_soc - difference_per_sortie/next(iter(env.scenario.relay_models.values())).usable_energy_kwh)
            >= next(iter(env.scenario.relay_models.values())).minimum_return_soc-1e-9
            for r in solution.relays),
        "note": "Same route and resource schedule replay; alternative mode was not reoptimized."}
    (out/"q3_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    vdir=out/"validation";vdir.mkdir(exist_ok=True)
    _book(vdir/"q3_validation.xlsx","独立核验",
          ["项目","值"],[[k,v] for k,v in summary["validation"].items()
                       if k!="issues"]+[["issue",x] for x in validation.issues])
    (vdir/"q3_validation.txt").write_text(
        f"Q3 independent validator: PASS\nNumerically verified continuous communication; "
        f"max verification interval = {validation.verification_max_step_s} s; "
        f"transition tolerance = {validation.transition_tolerance_s} s.\n"
        f"Detected outage duration = {validation.outage_s:.9f} s.\n"
        "This is a numerical check, not an analytic proof of all instants.\n",
        encoding="utf-8")
    logs=out/"logs";logs.mkdir(exist_ok=True)
    with (logs/"q3_search_history.csv").open("w",newline="",encoding="utf-8-sig") as stream:
        fields=["iteration","restart","operator","score","j1","j2_s",
                "estimated_blackout_s","routes"]
        writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader()
        writer.writerows(history)
    _official_template(root,solution,validation.intervals)
    _plots(root,env,solution,validation,pareto,history)
