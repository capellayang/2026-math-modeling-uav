"""Q4 auditable JSON, workbooks, official submission sheet and figures."""

from dataclasses import asdict
import json
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from relief_uav.data.models import Scenario
from relief_uav.validation.q4 import Q4Validation
from .model import Q4Solution, RESOURCE_KEYS


def _sheets(path: Path, sheets: dict[str, tuple[list[str], list[list]]]) -> None:
    book = Workbook()
    book.remove(book.active)
    for name, (headers, rows) in sheets.items():
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 30
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(64, max(16, len(header) * 1.4))
        for cell in ws[1]:
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="24445C")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    book.save(path)


def _chosen(result):
    return next((c for c in result.candidates if c.selected), None)


def _official(root: Path, solution: Q4Solution, out: Path) -> None:
    target = out / "结果提交_Q4.xlsx"
    shutil.copy2(root / "结果提交模板.xlsx", target)
    book = load_workbook(target)
    ws = book["Q4_分区配置"]
    expected = ["K（2或3）", "任务组编号", "服务区列表", "A型运输无人机数",
                "B型运输无人机数", "C型运输无人机数", "A型电池组数",
                "B型电池组数", "C型电池组数", "中继无人机数", "中继能源组件数"]
    if [ws.cell(1, i).value for i in range(1, 12)] != expected:
        raise ValueError("Official Q4 template headers changed")
    row = 2
    for result in solution.policy_results:
        if result.relay_policy != "strict_no_duplication" or not result.feasible:
            continue
        selected = _chosen(result)
        for group in selected.groups:
            r = group.minimum_recolored_requirement
            values = [result.k, group.group.group_id,
                      ",".join(group.group.service_areas),
                      *(r[key] for key in RESOURCE_KEYS)]
            for col, value in enumerate(values, 1):
                ws.cell(row, col).value = value
            row += 1
    book.save(target)


def _plots(out: Path, scenario: Scenario, solution: Q4Solution) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    figures = out / "figures"
    figures.mkdir(exist_ok=True)
    strict = {r.k: _chosen(r) for r in solution.policy_results
              if r.relay_policy == "strict_no_duplication" and r.feasible}
    colors = ("#2b6cb0", "#dd6b20", "#6b46c1")
    for k in (2, 3):
        fig, ax = plt.subplots(figsize=(8, 7))
        candidate = strict.get(k)
        if candidate:
            for i, group in enumerate(candidate.groups):
                nodes = [scenario.services[sid] for sid in group.group.service_areas]
                ax.scatter([n.longitude_deg for n in nodes],
                           [n.latitude_deg for n in nodes],
                           s=90, color=colors[i], label=group.group.group_id)
                for node in nodes:
                    ax.annotate(node.node_id, (node.longitude_deg, node.latitude_deg),
                                xytext=(3, 4), textcoords="offset points", fontsize=8)
        ax.scatter([scenario.dispatch.longitude_deg], [scenario.dispatch.latitude_deg],
                   marker="*", s=200, color="black", label="O01")
        ax.set_xlabel("Longitude (deg)"); ax.set_ylabel("Latitude (deg)")
        ax.set_title(f"Selected strict Q4 partition: K={k}"); ax.legend()
        fig.tight_layout(); fig.savefig(figures/f"q4_partition_map_k{k}.png", dpi=170)
        plt.close(fig)
    fig, ax = plt.subplots(figsize=(12, 5))
    keys = list(RESOURCE_KEYS)
    for i, k in enumerate((2, 3)):
        values = [strict[k].resource_vector[x] for x in keys] if k in strict else [0]*len(keys)
        ax.bar([j + (i - .5)*.35 for j in range(len(keys))], values,
               width=.35, label=f"K={k}", color=colors[i])
    ax.set_xticks(range(len(keys)), keys, rotation=35, ha="right")
    ax.set_ylabel("Minimum independent units"); ax.legend(); fig.tight_layout()
    fig.savefig(figures/"q4_resource_comparison.png", dpi=170); plt.close(fig)
    fig, ax = plt.subplots(figsize=(8, 5))
    labels = []
    values = []
    for k in (2, 3):
        for group in strict.get(k, ()).groups if k in strict else ():
            labels.append(group.group.group_id)
            values.append(group.total_workload_s)
    ax.bar(labels, values, color=colors[:2] * 3)
    ax.set_ylabel("Transport + relay workload (s)")
    ax.set_title("Project workload definition: charging/turnaround excluded")
    fig.tight_layout(); fig.savefig(figures/"q4_workload_balance.png", dpi=170)
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(12, 5))
    x = range(len(keys))
    ax.bar([i-.25 for i in x], [solution.stock[r] for r in keys], .25, label="Stock")
    for offset, k in ((0, 2), (.25, 3)):
        ax.bar([i+offset for i in x],
               [strict[k].resource_vector[r] for r in keys] if k in strict else [0]*len(keys),
               .25, label=f"K={k} required")
    ax.set_xticks(list(x), keys, rotation=35, ha="right")
    ax.set_ylabel("Units"); ax.legend(); fig.tight_layout()
    fig.savefig(figures/"q4_resource_gap.png", dpi=170); plt.close(fig)


def save_q4_outputs(root: Path, scenario: Scenario, solution: Q4Solution,
                    validation: Q4Validation) -> Path:
    if not validation.passed:
        raise ValueError("Q4 cannot be reported before validator PASS")
    out = root / "outputs/q4"
    out.mkdir(parents=True, exist_ok=True)
    payload = asdict(solution)
    (out / "q4_summary.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                                          encoding="utf-8")
    atoms = {}
    for name, components in (("transport_components", solution.transport_components),
                             ("strict_components", solution.strict_components)):
        atoms[name] = (["component_id", "service_areas", "transport_sorties",
                        "relay_sorties", "box_count", "cargo_mass_kg"],
                       [[c.component_id, ",".join(c.service_areas),
                         ",".join(c.transport_sorties), ",".join(c.relay_sorties),
                         c.box_count, c.cargo_mass_kg] for c in components])
    atoms["relay_bindings"] = (["relay_sortie_id", "transport_sorties", "service_areas"],
        [[b["relay_sortie_id"], ",".join(b["transport_sorties"]),
          ",".join(b["service_areas"])] for b in solution.relay_bindings])
    _sheets(out/"q4_atomic_components.xlsx", atoms)
    candidate_sheets = {}
    candidate_headers = ["candidate_id", "K", "relay_policy", "group_1_services",
        "group_2_services", "group_3_services", *RESOURCE_KEYS, "resource_scale",
        "partition_redundancy_total", "stock_gap_total", "workload_CV",
        "workload_range", "pareto", "selected"]
    for result in solution.policy_results:
        name = ("strict" if result.relay_policy == "strict_no_duplication" else
                "replicate") + f"_K{result.k}"
        rows = []
        for c in result.candidates:
            sites = [",".join(g.group.service_areas) for g in c.groups]
            rows.append([c.candidate_id, c.k, c.relay_policy,
                *(sites + [None]*(3-len(sites))),
                *(c.resource_vector[key] for key in RESOURCE_KEYS),
                c.resource_scale, c.partition_redundancy_total,
                c.stock_gap_total, c.workload_cv, c.workload_range,
                c.pareto, c.selected])
        candidate_sheets[name] = (candidate_headers, rows)
    for name in ("strict_K2", "strict_K3", "replicate_K2", "replicate_K3"):
        candidate_sheets.setdefault(name, (candidate_headers, []))
    _sheets(out/"q4_partition_candidates.xlsx", candidate_sheets)
    selected = [(r, _chosen(r)) for r in solution.policy_results if r.feasible]
    selected_headers = ["K", "relay_policy", "candidate_id", "group_id",
        "service_areas", "transport_sorties", "relay_tasks", "service_area_count",
        "box_count", "cargo_mass_kg", "transport_sortie_count", "relay_task_count",
        *RESOURCE_KEYS, "transport_workload_s", "relay_workload_s", "total_workload_s"]
    selected_rows = []
    for result, candidate in selected:
        for group in candidate.groups:
            selected_rows.append([result.k, result.relay_policy,
                candidate.candidate_id, group.group.group_id,
                ",".join(group.group.service_areas),
                ",".join(group.group.transport_sorties),
                ",".join(x["task_id"] for x in group.group.relay_tasks),
                group.service_area_count, group.box_count, group.cargo_mass_kg,
                group.transport_sortie_count, group.relay_task_count,
                *(group.minimum_recolored_requirement[key] for key in RESOURCE_KEYS),
                group.transport_workload_s, group.relay_workload_s,
                group.total_workload_s])
    selected_scores = [[r.k, r.relay_policy, c.candidate_id,
        c.canonical_group_signature, c.stock_gap_total, c.resource_scale,
        c.partition_redundancy_total, c.workload_cv, c.workload_range,
        len(r.candidates), len(r.pareto_candidate_ids)] for r, c in selected]
    _sheets(out/"q4_selected_partitions.xlsx", {
        "selected_scores": (["K", "relay_policy", "candidate_id",
            "canonical_group_signature", "stock_gap_total", "resource_scale",
            "partition_redundancy_total", "workload_CV", "workload_range",
            "candidate_count", "pareto_count"], selected_scores),
        "selected_groups": (selected_headers, selected_rows)})
    inherited_headers = ["inherited_" + key for key in RESOURCE_KEYS]
    requirement_rows = [row + [g.inherited_q3_id_count[key] for key in RESOURCE_KEYS]
        for (r, c) in selected for g, row in zip(c.groups,
            [x for x in selected_rows if x[0] == r.k and x[1] == r.relay_policy])]
    _sheets(out/"q4_resource_requirements.xlsx", {"requirements":
        (selected_headers + inherited_headers, requirement_rows)})
    coloring_sheets = {}
    coloring_headers = ["K", "policy", "group", "task_id", "resource_type",
        "internal_resource_id", "occupied_start_s", "occupied_end_s",
        "source_q3_resource_id", "source_q3_task_id"]
    for name, keys in (("运输无人机", {"A_UAV", "B_UAV", "C_UAV"}),
                       ("运输电池", {"A_battery", "B_battery", "C_battery"}),
                       ("中继无人机", {"relay_UAV"}),
                       ("中继组件", {"relay_component"})):
        rows = []
        for result, candidate in selected:
            for group in candidate.groups:
                for a in group.assignments:
                    if a.resource_type in keys:
                        rows.append([result.k, result.relay_policy,
                            a.group_id, a.task_id, a.resource_type,
                            a.internal_resource_id, a.occupied_start_s,
                            a.occupied_end_s, a.source_q3_resource_id,
                            a.source_task_id])
        coloring_sheets[name] = (coloring_headers, rows)
    _sheets(out/"q4_resource_coloring.xlsx", coloring_sheets)
    gap_headers = ["K", "policy", "candidate_id", "resource_type", "official_stock",
        "global_q3_minimum", "partition_required", "partition_overhead",
        "stock_surplus", "stock_gap", "peak_group", "peak_start_s",
        "peak_end_s", "peak_required", "critical_task_ids", "reason"]
    gap_rows = []
    for result, candidate in selected:
        for key in RESOURCE_KEYS:
            peak = max((p for g in candidate.groups for p in g.resource_peaks
                        if p.resource_type == key), key=lambda p: p.required)
            reason = ("Independent group allocations add across groups; fixed overlapping tasks"
                      + (" and full recharge" if key.endswith("battery") or
                           key == "relay_component" else "")
                      + (" / 300 s turnaround" if key == "relay_UAV" else ""))
            gap_rows.append([result.k, result.relay_policy, candidate.candidate_id,
                key, solution.stock[key], solution.global_q3_minimum[key],
                candidate.resource_vector[key], candidate.partition_redundancy[key],
                candidate.stock_surplus[key], candidate.stock_gap[key],
                peak.group_id, peak.start_s, peak.end_s, peak.required,
                ",".join(peak.critical_task_ids), reason])
    _sheets(out/"q4_resource_gap.xlsx", {"gaps_and_peaks": (gap_headers, gap_rows)})
    strict = {r.k: _chosen(r) for r in solution.policy_results
              if r.relay_policy == "strict_no_duplication" and r.feasible}
    comparison_rows = []
    fields = [*RESOURCE_KEYS, "resource_scale", "partition_redundancy_total",
              "stock_gap_total", "workload_cv", "workload_range"]
    for field in fields:
        def value(k):
            if k not in strict:
                return "INFEASIBLE"
            return (strict[k].resource_vector[field] if field in RESOURCE_KEYS
                    else getattr(strict[k], field))
        comparison_rows.append([field, value(2), value(3)])
    for k in (2, 3):
        for group in strict[k].groups if k in strict else ():
            comparison_rows.append([group.group.group_id + " services",
                ",".join(group.group.service_areas) if k == 2 else None,
                ",".join(group.group.service_areas) if k == 3 else None])
            for field in ("service_area_count", "transport_sortie_count",
                          "relay_task_count", "total_workload_s"):
                comparison_rows.append([group.group.group_id + " " + field,
                    getattr(group, field) if k == 2 else None,
                    getattr(group, field) if k == 3 else None])
    _sheets(out/"q4_k2_k3_comparison.xlsx", {"strict_K2_vs_K3":
        (["metric", "strict_K2", "strict_K3"], comparison_rows)})
    sensitivity_headers = ["K", "relay_policy", "atomic_component_count",
        "candidate_count", "selected_partition", "original_relay_task_count",
        "effective_relay_task_count", "added_relay_task_count", "relay_UAV_required",
        "component_required", "resource_scale", "partition_redundancy_total",
        "stock_gap_total", "workload_CV", "candidates_requiring_copies",
        "minimum_relay_tasks_all_candidates", "maximum_relay_tasks_all_candidates",
        "maximum_relay_UAV_all_candidates", "maximum_component_all_candidates",
        "maximum_resource_scale_all_candidates"]
    sensitivity_rows = []
    for result, candidate in selected:
        sensitivity_rows.append([result.k, result.relay_policy,
            result.atomic_component_count, len(result.candidates),
            candidate.canonical_group_signature, candidate.original_relay_task_count,
            candidate.effective_relay_task_count,
            candidate.effective_relay_task_count - candidate.original_relay_task_count,
            candidate.resource_vector["relay_UAV"],
            candidate.resource_vector["relay_component"], candidate.resource_scale,
            candidate.partition_redundancy_total, candidate.stock_gap_total,
            candidate.workload_cv,
            sum(c.effective_relay_task_count > c.original_relay_task_count
                for c in result.candidates),
            min(c.effective_relay_task_count for c in result.candidates),
            max(c.effective_relay_task_count for c in result.candidates),
            max(c.resource_vector["relay_UAV"] for c in result.candidates),
            max(c.resource_vector["relay_component"] for c in result.candidates),
            max(c.resource_scale for c in result.candidates)])
    _sheets(out/"q4_relay_policy_sensitivity.xlsx", {"relay_policy_sensitivity":
        (sensitivity_headers, sensitivity_rows)})
    _official(root, solution, out)
    _plots(out, scenario, solution)
    return out


def save_q4_validation(out: Path, validation: Q4Validation) -> None:
    folder = out / "validation"
    folder.mkdir(exist_ok=True)
    rows = [[key, value] for key, value in asdict(validation).items()
            if key != "issues"] + [["issue", issue] for issue in validation.issues]
    _sheets(folder/"q4_validation.xlsx", {"independent_validation":
        (["item", "value"], rows)})
    content = (
        f"Q4 independent validator: {'PASS' if validation.passed else 'FAIL'}\n"
        f"Q3 source SHA256: {validation.source_q3_sha256}\n"
        f"Checked {validation.checked_candidates} candidates, "
        f"{validation.checked_groups} groups, "
        f"{validation.checked_assignments} resource assignments.\n")
    if validation.issues:
        content += "\n".join(validation.issues) + "\n"
    (folder/"q4_validation.txt").write_text(content, encoding="utf-8")
