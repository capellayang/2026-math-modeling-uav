"""Python-only Q1 outputs, official-template copy and sensitivity figures."""

from dataclasses import asdict
import json
from pathlib import Path
import shutil
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from relief_uav.validation.q1 import Q1Validation
from .model import Q1Solution


def _format_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="24445C")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"


def _save_book(path: Path, title: str, header: list[str], rows: list[list], widths):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    _format_sheet(sheet, widths)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def save_q1_solution(root: Path, solution: Q1Solution, validation: Q1Validation) -> None:
    qdir = root / "outputs/q1"
    vdir = root / "outputs/validation"
    qdir.mkdir(parents=True, exist_ok=True)
    vdir.mkdir(parents=True, exist_ok=True)
    safe_rows = [[x.service_id, x.model_id, x.one_way_distance_m,
                  2 * x.one_way_distance_m, x.maximum_dem_m, x.cruise_altitude_m,
                  x.structural_payload_kg, x.safe_payload_kg, x.sortie_energy_kwh,
                  None if x.return_soc is None else x.return_soc * 100,
                  None if x.energy_limited is None else ("是" if x.energy_limited else "否")]
                 for x in solution.safe_payloads]
    _save_book(qdir / "q1_safe_payload_matrix.xlsx", "安全载荷",
               ["服务区编号", "机型编号", "单程水平距离（m）", "往返水平距离（m）",
                "最高DEM（m）", "巡航海拔（m）", "最大结构载荷（kg）",
                "最大安全载荷（kg）", "对应架次能耗（kWh）", "返航SOC（%）",
                "是否能量限制"], safe_rows,
               [15, 12, 23, 23, 18, 18, 22, 22, 25, 19, 17])

    trip_rows = [[t.trip_id, t.service_id, t.model_id, ", ".join(t.box_ids),
                  len(t.box_ids), t.mass_kg, t.volume_m3, t.outgoing_payload_kg,
                  t.outgoing_energy_kwh, t.returning_energy_kwh, t.total_energy_kwh,
                  t.preparation_s, t.outgoing_flight_s, t.handover_s,
                  t.returning_flight_s, t.operation_s, t.return_soc * 100]
                 for t in solution.trips]
    _save_book(qdir / "q1_plan.xlsx", "组批方案",
               ["架次编号", "服务区编号", "机型编号", "货箱编号列表", "箱数",
                "总质量（kg）", "总体积（m³）", "去程载荷（kg）",
                "去程能耗（kWh）", "回程能耗（kWh）", "总能耗（kWh）",
                "准备时间（s）", "去程飞行时间（s）", "交接时间（s）",
                "回程飞行时间（s）", "总作业时间（s）", "返航SOC（%）"],
               trip_rows, [16, 15, 13, 90, 10, 18, 18, 20, 22, 22, 20,
                           19, 22, 18, 22, 22, 18])

    summary = {
        "model_status": "project modeling assumptions for horizontal/climb energy; not official formulas",
        "objective_mode": solution.objective_mode,
        "safety_margin_soc": solution.safety_margin_soc,
        "total_trips": solution.total_trips,
        "total_energy_kwh": solution.total_energy_kwh,
        "total_operation_s": solution.total_operation_s,
        "by_service": {s: {"trip_count": sum(t.service_id == s for t in solution.trips),
                           "models": [t.model_id for t in solution.trips if t.service_id == s]}
                       for s in sorted({t.service_id for t in solution.trips})},
        "safe_payloads": [asdict(x) for x in solution.safe_payloads],
        "trips": [asdict(t) for t in solution.trips],
        "validation_passed": validation.passed,
    }
    (qdir / "q1_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    validation_rows = [["总体", "通过" if validation.passed else "失败",
                        f"{validation.checked_trips}架次，{validation.delivered_unique_boxes}/{validation.expected_boxes}箱唯一交付"]]
    validation_rows.extend([[i.trip_id, i.check, i.detail] for i in validation.issues])
    _save_book(vdir / "q1_validation.xlsx", "独立核验", ["架次", "状态或检查", "详情"],
               validation_rows, [20, 30, 100])
    text = [f"Q1 independent validator: {'PASS' if validation.passed else 'FAIL'}",
            f"Trips checked: {validation.checked_trips}",
            f"Unique boxes: {validation.delivered_unique_boxes}/{validation.expected_boxes}",
            "Checks: exactly-once box coverage, model/service, structural mass, volume,",
            "common-physics energy, return SOC, operation time, detail/summary totals."]
    text.extend(f"{i.trip_id} | {i.check} | {i.detail}" for i in validation.issues)
    (vdir / "q1_validation.txt").write_text("\n".join(text) + "\n", encoding="utf-8")

    # Copy the official workbook first; populate only its existing Q1 sheet.
    template = root / "结果提交模板.xlsx"
    target = qdir / "结果提交_Q1.xlsx"
    shutil.copy2(template, target)
    workbook = load_workbook(target)
    sheet = workbook["Q1_单点组批"]
    expected_headers = ["架次编号", "服务区编号", "机型编号", "货箱编号列表",
                        "总质量（kg）", "总体积（m³）", "往返时间（s）",
                        "架次能耗（kWh）", "返航SOC（%）"]
    if [sheet.cell(1, c).value for c in range(1, 10)] != expected_headers:
        raise ValueError("Official Q1 template headers changed")
    from copy import copy
    for row_number, trip in enumerate(solution.trips, 2):
        values = [trip.trip_id, trip.service_id, trip.model_id, ", ".join(trip.box_ids),
                  trip.mass_kg, trip.volume_m3,
                  trip.outgoing_flight_s + trip.returning_flight_s,
                  trip.total_energy_kwh, trip.return_soc * 100]
        for col_number, value in enumerate(values, 1):
            cell = sheet.cell(row_number, col_number)
            if row_number > 2 and sheet.cell(2, col_number).has_style:
                cell._style = copy(sheet.cell(2, col_number)._style)
            cell.value = value
    workbook.save(target)


def save_sensitivity(root: Path, solutions: list[Q1Solution]) -> None:
    qdir = root / "outputs/q1"
    qdir.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary.append(["返航安全余量（%）", "总架次数", "总运输能耗（kWh）", "累计作业时间（s）"])
    payload = workbook.create_sheet("安全载荷")
    payload.append(["返航安全余量（%）", "服务区编号", "机型编号", "最大安全载荷（kg）",
                    "对应架次能耗（kWh）", "返航SOC（%）", "能量限制"])
    batches = workbook.create_sheet("逐档组批")
    batches.append(["返航安全余量（%）", "架次编号", "服务区编号", "机型编号",
                    "货箱编号列表", "能耗（kWh）", "作业时间（s）", "返航SOC（%）"])
    for s in solutions:
        pct = round(s.safety_margin_soc * 100, 8)
        summary.append([pct, s.total_trips, s.total_energy_kwh, s.total_operation_s])
        for x in s.safe_payloads:
            payload.append([pct, x.service_id, x.model_id, x.safe_payload_kg,
                            x.sortie_energy_kwh, None if x.return_soc is None else 100 * x.return_soc,
                            None if x.energy_limited is None else ("是" if x.energy_limited else "否")])
        for t in s.trips:
            batches.append([pct, t.trip_id, t.service_id, t.model_id,
                            ", ".join(t.box_ids), t.total_energy_kwh,
                            t.operation_s, t.return_soc * 100])
    for sheet, widths in [(summary, [24, 16, 26, 26]),
                          (payload, [24, 16, 14, 24, 25, 20, 16]),
                          (batches, [24, 16, 16, 14, 90, 23, 23, 20])]:
        _format_sheet(sheet, widths)
    workbook.save(qdir / "q1_sensitivity.xlsx")
    _save_sensitivity_plots(root, solutions)


def _save_sensitivity_plots(root: Path, solutions: list[Q1Solution]):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    font_candidates = [p for p in font_manager.findSystemFonts() if "msyh" in p.lower()]
    if font_candidates:
        font_manager.fontManager.addfont(font_candidates[0])
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=font_candidates[0]).get_name()
    figure_dir = root / "outputs/figures"
    figure_dir.mkdir(parents=True, exist_ok=True)
    x = [s.safety_margin_soc * 100 for s in solutions]
    fig, axes = plt.subplots(1, 3, figsize=(15, 5.5))
    service_labels = sorted({z.service_id for s in solutions for z in s.safe_payloads})
    for axis, model_id in zip(axes, ["A", "B", "C"]):
        for service in service_labels:
            values = [next(z.safe_payload_kg for z in s.safe_payloads
                           if z.service_id == service and z.model_id == model_id)
                      for s in solutions]
            axis.plot(x, values, linewidth=1.1, alpha=0.8, label=service)
        axis.set_title(f"{model_id}型：各服务区")
        axis.set_xlabel("返航安全余量（%）")
        axis.set_ylabel("最大安全载荷（kg）")
        axis.grid(alpha=0.25)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=8, frameon=False, fontsize=8)
    fig.subplots_adjust(bottom=0.27, left=0.06, right=0.99, wspace=0.31)
    fig.savefig(figure_dir / "q1_safe_payload_sensitivity.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 3, figsize=(12, 4), constrained_layout=True)
    series = [("总架次数", [s.total_trips for s in solutions]),
              ("总运输能耗（kWh）", [s.total_energy_kwh for s in solutions]),
              ("累计作业时间（s）", [s.total_operation_s for s in solutions])]
    for axis, (label, values) in zip(axes, series):
        axis.plot(x, values, marker="o", color="#1F618D")
        axis.set_xlabel("返航安全余量（%）")
        axis.set_ylabel(label)
        axis.grid(alpha=0.25)
    from matplotlib.ticker import MaxNLocator
    axes[0].yaxis.set_major_locator(MaxNLocator(integer=True))
    fig.savefig(figure_dir / "q1_objective_sensitivity.png", dpi=180)
    plt.close(fig)


def replot_saved_sensitivity(root: Path) -> None:
    """Restyle figures from saved results without rerunning any Q1 optimization."""
    book = load_workbook(root / "outputs/q1/q1_sensitivity.xlsx", read_only=True,
                         data_only=True)
    summaries = {row[0]: row[1:] for row in list(book["汇总"].values)[1:]}
    payloads = {margin: [] for margin in summaries}
    for row in list(book["安全载荷"].values)[1:]:
        payloads[row[0]].append(SimpleNamespace(service_id=row[1], model_id=row[2],
                                                safe_payload_kg=row[3]))
    solutions = [SimpleNamespace(safety_margin_soc=margin / 100,
                                 total_trips=summaries[margin][0],
                                 total_energy_kwh=summaries[margin][1],
                                 total_operation_s=summaries[margin][2],
                                 safe_payloads=payloads[margin])
                 for margin in sorted(summaries)]
    _save_sensitivity_plots(root, solutions)
