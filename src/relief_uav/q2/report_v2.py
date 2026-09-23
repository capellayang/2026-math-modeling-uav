"""Versioned Q2-v2 outputs; the Q2-v1 directory is never written here."""

from collections import Counter
from dataclasses import asdict
import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from relief_uav.validation.q2 import Q2Validation
from .model import Q2Solution
from .objectives import epsilon_feasible, primary_objectives, secondary_objectives
from .report import load_q2_solution, save_q2_outputs
from .search_v2 import V2Run


def _comparison_rows(run: V2Run) -> list[list]:
    baseline = run.baseline.objective
    rows = [("Q2-v1 基准", None, run.baseline, "FEASIBLE", run.baseline.search_seconds)]
    rows.extend((f"Q2-v2 ε={epsilon:.0%}", epsilon, candidate.solution,
                 candidate.metadata.cp_sat_status, run.statistics.elapsed_s)
                for epsilon, candidate in sorted(run.epsilon_results.items()))
    rows.append(("Q2-v2 最终选择", run.config.tardiness_slack,
                 run.selected.solution, run.selected.metadata.cp_sat_status,
                 run.statistics.elapsed_s))
    output = []
    for label, epsilon, solution, status, runtime in rows:
        obj = solution.objective
        output.append([label, epsilon, obj.weighted_tardiness,
                       obj.normalized_weighted_tardiness, obj.makespan_s,
                       obj.total_energy_kwh, obj.sortie_count, runtime, status,
                       (obj.weighted_tardiness / baseline.weighted_tardiness - 1) * 100
                       if baseline.weighted_tardiness else None,
                       (obj.makespan_s / baseline.makespan_s - 1) * 100,
                       (obj.total_energy_kwh / baseline.total_energy_kwh - 1) * 100,
                       (obj.sortie_count / baseline.sortie_count - 1) * 100])
    return output


def _save_comparison(path: Path, run: V2Run) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "v1-v2 对比"
    sheet.append(["方案", "ε", "J1加权逾期", "J1归一化", "J2最晚返航（s）",
                  "J3能耗（kWh）", "J4架次", "运行时间（s）", "CP-SAT状态",
                  "J1相对v1变化（%）", "J2相对v1变化（%）",
                  "J3相对v1变化（%）", "J4相对v1变化（%）"])
    for row in _comparison_rows(run):
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column in range(1, 14):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(column)].width = 24
    book.save(path)


def _save_pareto(path: Path, run: V2Run) -> None:
    book = Workbook()
    front = book.active
    front.title = "J1-J2 Pareto前沿"
    all_sheet = book.create_sheet("全部CP候选")
    epsilon_sheet = book.create_sheet("ε扫描")
    headers = ["候选编号", "seed", "ε", "J1加权逾期", "J1归一化", "J2最晚返航（s）",
               "J3能耗（kWh）", "J4架次", "CP-SAT状态", "是否选中", "路线摘要"]
    for sheet in (front, all_sheet):
        sheet.append(headers)
        sheet.freeze_panes = "A2"
    def row(candidate):
        m = candidate.metadata
        return [m.candidate_id, m.seed, m.epsilon_level, m.weighted_tardiness,
                m.normalized_weighted_tardiness, m.makespan_s,
                m.total_energy_kwh, m.sortie_count, m.cp_sat_status,
                "是" if m.candidate_id == run.selected.metadata.candidate_id else "否",
                m.route_signature]
    for candidate in run.pareto:
        front.append(row(candidate))
    for candidate in run.all_cp_candidates:
        all_sheet.append(row(candidate))
    epsilon_sheet.append(["ε", "候选编号", "J1加权逾期", "J1归一化",
                          "J2最晚返航（s）", "J3能耗（kWh）", "J4架次",
                          "CP-SAT状态"])
    for epsilon, candidate in sorted(run.epsilon_results.items()):
        o = candidate.solution.objective
        epsilon_sheet.append([epsilon, candidate.metadata.candidate_id,
                              o.weighted_tardiness, o.normalized_weighted_tardiness,
                              o.makespan_s, o.total_energy_kwh, o.sortie_count,
                              candidate.metadata.cp_sat_status])
    for sheet in book:
        from openpyxl.utils import get_column_letter
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 23 if index < 11 else 100
    book.save(path)


def _save_history(path: Path, run: V2Run) -> None:
    headers = ["restart", "seed", "iteration", "global_iteration", "destroy_operator",
               "repair_operator", "destroy_size", "accepted", "accept_reason",
               "epsilon", "J1", "J2", "J3", "J4", "pareto_archive_size",
               "best_J1", "best_J2", "temperature", "destroy_weight", "repair_weight"]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(run.history)


def _save_figures(folder: Path, run: V2Run) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    folder.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(8, 5))
    front = run.pareto
    ax.scatter([row.solution.objective.normalized_weighted_tardiness for row in front],
               [row.solution.objective.makespan_s for row in front],
               color="#264653", s=40, label="Q2-v2 Pareto")
    for row in front:
        ax.annotate(row.metadata.candidate_id,
                    (row.solution.objective.normalized_weighted_tardiness,
                     row.solution.objective.makespan_s), fontsize=7)
    ax.scatter([run.baseline.objective.normalized_weighted_tardiness],
               [run.baseline.objective.makespan_s], marker="x", s=90,
               color="#d62828", label="Q2-v1")
    ax.scatter([run.selected.solution.objective.normalized_weighted_tardiness],
               [run.selected.solution.objective.makespan_s], marker="*", s=190,
               color="#f4a261", label="Selected")
    ax.set_xlabel("Normalized weighted tardiness")
    ax.set_ylabel("Makespan (s)")
    ax.set_title("Q2-v2 timeliness / makespan Pareto front")
    ax.grid(alpha=0.2)
    ax.legend()
    fig.tight_layout()
    fig.savefig(folder / "q2_timeliness_makespan_pareto.png", dpi=180)
    plt.close(fig)
    levels = sorted(run.epsilon_results)
    objectives = [run.epsilon_results[level].solution.objective for level in levels]
    fig, axes = plt.subplots(2, 2, figsize=(10, 7), constrained_layout=True)
    series = [("Weighted tardiness", [o.weighted_tardiness for o in objectives]),
              ("Makespan (s)", [o.makespan_s for o in objectives]),
              ("Transport energy (kWh)", [o.total_energy_kwh for o in objectives]),
              ("Sortie count", [o.sortie_count for o in objectives])]
    for axis, (label, values) in zip(axes.flat, series):
        axis.plot([100 * x for x in levels], values, marker="o", color="#287271")
        axis.set_xlabel("Relative epsilon (%)")
        axis.set_ylabel(label)
        axis.grid(alpha=0.2)
    fig.savefig(folder / "q2_epsilon_sensitivity.png", dpi=180)
    plt.close(fig)
    weights = {f"D: {name}": weight for name, weight in run.statistics.destroy_weights.items()}
    weights.update({f"R: {name}": weight for name, weight in run.statistics.repair_weights.items()})
    fig, ax = plt.subplots(figsize=(10, 5))
    labels = list(weights)
    ax.barh(labels, [weights[name] for name in labels], color="#457b9d")
    ax.set_xlabel("Final adaptive weight")
    ax.set_title("Q2-v2 ALNS operator weights")
    fig.tight_layout()
    fig.savefig(folder / "q2_operator_weights.png", dpi=180)
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(9, 4))
    ax.plot([row["global_iteration"] for row in run.history],
            [row["best_J1"] for row in run.history], color="#1d3557")
    ax.set_xlabel("ALNS iteration")
    ax.set_ylabel("Best fast-schedule J1")
    ax.grid(alpha=0.2)
    fig.tight_layout()
    fig.savefig(folder / "q2_objective_history.png", dpi=180)
    plt.close(fig)


def save_v2_outputs(root: Path, run: V2Run, final_solution: Q2Solution,
                    validation: Q2Validation) -> None:
    if not validation.passed:
        raise ValueError("Q2-v2 final validator must PASS before writing outputs")
    if run.selected.metadata.candidate_id not in {c.metadata.candidate_id for c in run.pareto}:
        raise ValueError("Selected Q2-v2 candidate is absent from Pareto front")
    out = root / "outputs/q2_v2"
    validation_dir = out / "validation"
    figure_dir = out / "figures"
    legacy_pareto = tuple((candidate.specs, candidate.solution,
                           candidate.metadata.cp_sat_status) for candidate in run.pareto)
    save_q2_outputs(root, final_solution, validation, legacy_pareto, (), (),
                    run.selected.metadata.cp_sat_status,
                    output_dir=out, validation_dir=validation_dir,
                    figure_dir=figure_dir)
    _save_pareto(out / "q2_pareto.xlsx", run)
    _save_comparison(out / "q2_algorithm_comparison.xlsx", run)
    _save_history(out / "logs/q2_search_history.csv", run)
    _save_figures(figure_dir, run)
    summary_path = out / "q2_summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary.update({
        "algorithm_version": "Q2-v2",
        "selection_method": run.config.selection_method,
        "epsilon_level": run.config.tardiness_slack,
        "candidate_epsilon_level": run.selected.metadata.epsilon_level,
        "primary_objectives": list(primary_objectives(final_solution)),
        "secondary_objectives": list(secondary_objectives(final_solution)),
        "selected_candidate_id": run.selected.metadata.candidate_id,
        "j1_best_found": run.j1_best_found,
        "epsilon_results": {str(level): candidate.metadata.candidate_id
                            for level, candidate in run.epsilon_results.items()},
        "pareto_candidates": [asdict(candidate.metadata) for candidate in run.pareto],
        "all_cp_candidate_count": len(run.all_cp_candidates),
        "search_statistics": asdict(run.statistics),
        "algorithm_config": asdict(run.config),
        "baseline_v1_summary": "baseline_v1/q2_summary.json",
    })
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2),
                            encoding="utf-8")


def benchmark_saved_v2(root: Path) -> dict:
    """Read saved results and compare automatically; no manually copied metrics."""
    out = root / "outputs/q2_v2"
    baseline = load_q2_solution(out / "baseline_v1/q2_summary.json")
    current = load_q2_solution(out / "q2_summary.json")
    with (out / "q2_summary.json").open(encoding="utf-8") as stream:
        metadata = json.load(stream)
    from openpyxl import load_workbook as read_book
    book = read_book(out / "q2_algorithm_comparison.xlsx", read_only=True, data_only=True)
    rows = list(book.active.values)
    book.close()
    result = {
        "baseline": asdict(baseline.objective),
        "selected": asdict(current.objective),
        "selection_method": metadata["selection_method"],
        "epsilon_level": metadata["epsilon_level"],
        "comparison_headers": rows[0],
        "comparison_rows": rows[1:],
    }
    (out / "q2_benchmark.json").write_text(json.dumps(result, ensure_ascii=False, indent=2),
                                            encoding="utf-8")
    return result
