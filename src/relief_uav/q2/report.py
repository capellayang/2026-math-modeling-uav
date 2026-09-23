"""Q2 supplementary workbooks, official-template copy, JSON and charts."""

from collections import Counter
from dataclasses import asdict
import json
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from relief_uav.validation.q2 import Q2Validation
from .model import (Q2BatteryUse, Q2BoxDelivery, Q2ChargeEvent, Q2LegTimeline,
                    Q2Objective, Q2Solution, Q2Sortie, Q2SortieSpec, Q2StopTimeline,
                    StopAssignment)


def _book(path: Path, title: str, headers: list[str], rows: list[list]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for i, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = min(65, max(17, len(header) * 1.8))
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="24445C")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def load_q2_solution(path: Path) -> Q2Solution:
    data = json.loads(path.read_text(encoding="utf-8"))
    sorties = []
    for raw in data["sorties"]:
        spec = raw["spec"]
        spec = Q2SortieSpec(spec["sortie_id"], spec["model_id"], tuple(spec["route"]),
                            tuple(StopAssignment(x["service_id"], tuple(x["box_ids"]))
                                  for x in spec["deliveries"]))
        sorties.append(Q2Sortie(spec, raw["drone_id"], raw["battery_id"],
                                raw["preparation_start_s"], raw["takeoff_time_s"],
                                tuple(Q2LegTimeline(**x) for x in raw["legs"]),
                                tuple(Q2StopTimeline(x["service_id"], tuple(x["box_ids"]),
                                                     x["arrival_time_s"], x["handover_start_s"],
                                                     x["handover_end_s"], x["delivery_complete_time_s"])
                                      for x in raw["stops"]), raw["return_o01_time_s"],
                                raw["loaded_mass_kg"], raw["loaded_volume_m3"],
                                raw["energy_kwh"], raw["return_soc"]))
    if data.get("model_counts") != dict(Counter(s.spec.model_id for s in sorties)):
        raise ValueError("Q2 summary model counts disagree with sortie detail")
    if data.get("used_drone_ids") != sorted({s.drone_id for s in sorties}):
        raise ValueError("Q2 summary drone IDs disagree with sortie detail")
    if data.get("used_batteries_by_model") != {
            m: len({s.battery_id for s in sorties if s.spec.model_id == m}) for m in "ABC"}:
        raise ValueError("Q2 summary battery counts disagree with sortie detail")
    return Q2Solution(data["objective_mode"], data["seed"], tuple(sorties),
                      tuple(Q2BoxDelivery(**x) for x in data["box_deliveries"]),
                      tuple(Q2BatteryUse(**x) for x in data["battery_uses"]),
                      tuple(Q2ChargeEvent(**x) for x in data["charge_events"]),
                      Q2Objective(**data["objective"]), data["search_seconds"],
                      data["pareto_count"], data["energy_model_status"])


def save_q2_outputs(root: Path, solution: Q2Solution, validation: Q2Validation,
                    pareto: tuple, search_pareto: tuple, history: tuple,
                    cp_status: str, *, output_dir: Path | None = None,
                    validation_dir: Path | None = None,
                    figure_dir: Path | None = None) -> None:
    if not validation.passed:
        raise ValueError("Cannot save Q2 final outputs before independent validation PASS")
    out = output_dir or root / "outputs/q2"
    out.mkdir(parents=True, exist_ok=True)
    _book(out / "q2_transport_sorties.xlsx", "架次",
          ["架次编号", "机型编号", "无人机编号", "电池编号", "路线", "逐站货箱",
           "准备开始（s）", "起飞（s）", "返回O01（s）", "质量（kg）", "体积（m³）",
           "能耗（kWh）", "返航SOC（%）"],
          [[s.spec.sortie_id, s.spec.model_id, s.drone_id, s.battery_id,
            "→".join(s.spec.route), "; ".join(f"{stop.service_id}:{','.join(stop.box_ids)}"
                                              for stop in s.spec.deliveries),
            s.preparation_start_s, s.takeoff_time_s, s.return_o01_time_s,
            s.loaded_mass_kg, s.loaded_volume_m3, s.energy_kwh, 100 * s.return_soc]
           for s in solution.sorties])
    _book(out / "q2_box_deliveries.xlsx", "逐箱交付",
          ["货箱编号", "架次编号", "服务区编号", "交付完成（s）", "期望送达（s）",
           "首批截止（s）", "优先系数", "加权逾期（系数·s）"],
          [[d.box_id, d.sortie_id, d.service_id, d.delivery_complete_time_s,
            d.desired_delivery_s, d.first_batch_deadline_s, d.emergency_priority,
            d.weighted_tardiness] for d in sorted(solution.box_deliveries, key=lambda x: x.box_id)])
    leg_rows = []
    stop_rows = []
    for s in solution.sorties:
        for index, leg in enumerate(s.legs, 1):
            leg_rows.append([s.spec.sortie_id, s.drone_id, index, leg.start_id, leg.end_id,
                             leg.payload_kg, leg.distance_m, leg.maximum_dem_m,
                             leg.cruise_altitude_m, leg.climb_m, leg.descent_m,
                             leg.equivalent_range_m, leg.horizontal_energy_kwh,
                             leg.climb_energy_kwh, leg.climb_start_s, leg.climb_end_s,
                             leg.cruise_start_s, leg.cruise_end_s,
                             leg.descent_start_s, leg.descent_end_s])
        for stop in s.stops:
            stop_rows.append([s.spec.sortie_id, s.drone_id, stop.service_id,
                              ",".join(stop.box_ids), stop.arrival_time_s,
                              stop.handover_start_s, stop.handover_end_s,
                              stop.delivery_complete_time_s])
    _book(out / "q2_drone_timeline.xlsx", "航段时间轴",
          ["架次编号", "无人机编号", "段序号", "起点", "终点", "剩余载荷（kg）",
           "水平距离（m）", "最高DEM（m）", "巡航海拔（m）", "爬升（m）", "下降（m）",
           "等效航程（m）", "水平能耗（kWh）", "爬升能耗（kWh）",
           "爬升开始（s）", "爬升结束（s）", "巡航开始（s）", "巡航结束（s）",
           "下降开始（s）", "下降结束（s）"], leg_rows)
    # Stop timeline is also preserved in JSON for Q3. A separate sheet provides
    # a human-readable view without hiding any flight stages.
    book = load_workbook(out / "q2_drone_timeline.xlsx")
    ws = book.create_sheet("交接时间轴")
    ws.append(["架次编号", "无人机编号", "服务区", "货箱", "到达（s）", "交接开始（s）",
               "交接结束（s）", "交付完成（s）"])
    for row in stop_rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i in range(1, 9):
        ws.column_dimensions[get_column_letter(i)].width = 22 if i != 4 else 70
    book.save(out / "q2_drone_timeline.xlsx")
    _book(out / "q2_battery_timeline.xlsx", "电池占用与充电",
          ["电池编号（内部）", "机型", "架次", "分配开始（s）", "返航（s）",
           "开始SOC（%）", "返航SOC（%）", "充电开始（s）", "充电结束（s）"],
          [[u.battery_id, u.model_id, u.sortie_id, u.use_start_s, u.use_end_s,
            100 * u.start_soc, 100 * u.return_soc, c.charge_start_s, c.charge_end_s]
           for u, c in zip(solution.battery_uses, solution.charge_events)])
    summary = asdict(solution)
    summary["cp_sat_status"] = cp_status
    summary["search_pareto_count"] = len(search_pareto)
    summary["validation_passed"] = validation.passed
    summary["model_counts"] = dict(Counter(s.spec.model_id for s in solution.sorties))
    summary["used_drone_ids"] = sorted({s.drone_id for s in solution.sorties})
    summary["used_batteries_by_model"] = {m: len({s.battery_id for s in solution.sorties
                                                   if s.spec.model_id == m}) for m in "ABC"}
    (out / "q2_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2),
                                          encoding="utf-8")
    _book(out / "q2_pareto.xlsx", "非支配候选",
          ["候选编号", "加权逾期（系数·s）", "归一化加权逾期", "最晚返航（s）",
           "运输能耗（kWh）", "架次数", "CP-SAT状态", "路线与机型"],
          [[i, row[1].objective.weighted_tardiness,
            row[1].objective.normalized_weighted_tardiness,
            row[1].objective.makespan_s, row[1].objective.total_energy_kwh,
            row[1].objective.sortie_count, row[2],
            "; ".join(f"{s.model_id}:{'→'.join(s.route)}" for s in row[0])]
           for i, row in enumerate(pareto, 1)])
    archive_book = load_workbook(out / "q2_pareto.xlsx")
    archive_sheet = archive_book.create_sheet("快速搜索非支配档案")
    archive_sheet.append(["候选编号", "加权逾期（系数·s）", "归一化加权逾期",
                          "最晚返航（s）", "运输能耗（kWh）", "架次数", "路线与机型"])
    for i, (route, estimated) in enumerate(search_pareto, 1):
        objective = estimated.objective
        archive_sheet.append([i, objective.weighted_tardiness,
                              objective.normalized_weighted_tardiness,
                              objective.makespan_s, objective.total_energy_kwh,
                              objective.sortie_count,
                              "; ".join(f"{s.model_id}:{'→'.join(s.route)}" for s in route)])
    archive_sheet.freeze_panes = "A2"
    for column in range(1, 8):
        archive_sheet.column_dimensions[get_column_letter(column)].width = 25 if column < 7 else 100
    archive_book.save(out / "q2_pareto.xlsx")
    vdir = validation_dir or root / "outputs/validation"
    _book(vdir / "q2_validation.xlsx", "独立核验",
          ["对象", "检查", "详情"],
          [["总体", "PASS", f"{validation.delivered_unique_boxes}/{validation.expected_boxes}箱，{validation.checked_sorties}架次"]] +
          [[x.scope, x.check, x.detail] for x in validation.issues])
    (vdir / "q2_validation.txt").write_text(
        f"Q2 independent validator: PASS\n{validation.delivered_unique_boxes}/{validation.expected_boxes} boxes, {validation.checked_sorties} sorties\n"
        "Rebuilt from source boxes, route/model, drone/battery IDs and preparation starts.\n"
        "Checks: physical legs, time events, SOC/charge, deadlines, resource overlaps and totals.\n",
        encoding="utf-8")
    _save_template(root, solution, out)
    _plots(root, solution, history, figure_dir)


def _save_template(root: Path, solution: Q2Solution,
                   output_dir: Path | None = None) -> None:
    from copy import copy
    target = (output_dir or root / "outputs/q2") / "结果提交_Q2.xlsx"
    shutil.copy2(root / "结果提交模板.xlsx", target)
    workbook = load_workbook(target)
    a = workbook["Q2_运输架次"]
    b = workbook["Q2_逐箱交付"]
    if [a.cell(1, col).value for col in range(1, 9)] != [
            "架次编号", "无人机编号", "机型编号", "电池编号", "开始时刻（s）",
            "访问服务区顺序", "返回O01时刻（s）", "架次能耗（kWh）"]:
        raise ValueError("Official Q2 sortie template headers changed")
    if [b.cell(1, col).value for col in range(1, 5)] != [
            "货箱编号", "架次编号", "服务区编号", "交付完成时刻（s）"]:
        raise ValueError("Official Q2 box template headers changed")
    def put(sheet, row_number, values):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column)
            if row_number > 2 and sheet.cell(2, column).has_style:
                cell._style = copy(sheet.cell(2, column)._style)
            cell.value = value

    for row_number, s in enumerate(solution.sorties, 2):
        put(a, row_number, [s.spec.sortie_id, s.drone_id, s.spec.model_id, s.battery_id,
                            s.preparation_start_s, "→".join(s.spec.route[1:-1]),
                            s.return_o01_time_s, s.energy_kwh])
    for row_number, d in enumerate(sorted(solution.box_deliveries, key=lambda x: x.box_id), 2):
        put(b, row_number, [d.box_id, d.sortie_id, d.service_id,
                            d.delivery_complete_time_s])
    workbook.save(target)


def _plots(root: Path, solution: Q2Solution, history: tuple,
           figure_dir: Path | None = None) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    figdir = figure_dir or root / "outputs/figures"
    figdir.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(12, 5))
    drone_ids = sorted({s.drone_id for s in solution.sorties})
    for s in solution.sorties:
        ax.barh(drone_ids.index(s.drone_id), s.return_o01_time_s - s.preparation_start_s,
                left=s.preparation_start_s, height=0.6,
                color={"A": "#2a9d8f", "B": "#e9c46a", "C": "#e76f51"}[s.spec.model_id])
    ax.set_yticks(range(len(drone_ids)), drone_ids)
    ax.set_xlabel("Seconds from mission start")
    ax.set_title("Transport drone assignments (preparation through return)")
    ax.grid(axis="x", alpha=0.2)
    fig.tight_layout()
    fig.savefig(figdir / "q2_drone_gantt.png", dpi=170)
    plt.close(fig)
    deliveries = sorted(solution.box_deliveries, key=lambda d: d.delivery_complete_time_s)
    fig, ax = plt.subplots(figsize=(12, 4))
    ax.scatter(range(len(deliveries)), [d.delivery_complete_time_s for d in deliveries],
               c=["#d62828" if d.weighted_tardiness > 0 else "#2a9d8f" for d in deliveries], s=22)
    ax.plot(range(len(deliveries)), [d.desired_delivery_s for d in deliveries],
            color="#264653", linewidth=1, alpha=0.6)
    ax.set_xlabel("Boxes ordered by delivery time")
    ax.set_ylabel("Seconds")
    ax.set_title("Box delivery and desired times")
    ax.grid(alpha=0.2)
    fig.tight_layout()
    fig.savefig(figdir / "q2_delivery_tardiness.png", dpi=170)
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(10, 4))
    xs = [r.iteration for r in history if r.best_weighted_tardiness is not None]
    ys = [r.best_weighted_tardiness for r in history if r.best_weighted_tardiness is not None]
    ax.plot(xs, ys, color="#1d3557")
    ax.set_xlabel("Search iteration")
    ax.set_ylabel("Best fast-schedule weighted tardiness")
    ax.set_title("Q2 route-search progress")
    ax.grid(alpha=0.2)
    fig.tight_layout()
    fig.savefig(figdir / "q2_objective_history.png", dpi=170)
    plt.close(fig)
